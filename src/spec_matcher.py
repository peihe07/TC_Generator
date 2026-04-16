"""Spec reference matcher (RULES.md §2.4) — Layer 1: PDM code exact match."""
import re

from openpyxl import load_workbook


# Regex for all PDM code patterns
PDM_PATTERN = re.compile(
    r"\b(PDM\d+\.?\d*|PDMS\d+\.?\d*|MPDM\d+\.?\d*|TD\d+\.?\d*"
    r"|DNDS\d+\.?\d*|PDEE\d+\.?\d*|APAC\d+\.?\d*|PSR\d+\.?\d*)\b"
)


def extract_pdm_codes(text: str) -> list[str]:
    """Extract all PDM-family codes from text."""
    return PDM_PATTERN.findall(text)


def build_spec_index(sys1_xlsx_path: str) -> dict:
    """
    Parse SYS1 spec xlsx (Basic Report sheet) and build PDM code index.

    Returns:
        dict mapping PDM code -> {nrl_id, outline, source_id, description}
    """
    wb = load_workbook(sys1_xlsx_path, read_only=True, data_only=True)
    ws = wb["Basic Report"]

    index = {}
    row_num = 2  # Skip header row
    while True:
        nrl_id = ws.cell(row=row_num, column=1).value
        if not nrl_id:
            break

        outline = ws.cell(row=row_num, column=3).value
        description = ws.cell(row=row_num, column=4).value or ""
        source_id = ws.cell(row=row_num, column=5).value or ""

        # Extract PDM codes from description
        codes = extract_pdm_codes(description)
        for code in codes:
            index[code] = {
                "nrl_id": nrl_id,
                "outline": outline,
                "source_id": source_id,
                "description": description,
            }

        row_num += 1

    wb.close()
    return index


def match_spec_references(
    rows: list[dict],
    spec_index: dict,
) -> list[dict]:
    """
    Match Test Items to spec references using PDM code exact match.

    Args:
        rows: List of row dicts with 'req_id' and 'test_item' keys
        spec_index: PDM code index from build_spec_index()

    Returns:
        List of dicts with added 'spec_reference' and 'match_type' keys
    """
    results = []
    for row in rows:
        codes = extract_pdm_codes(row["test_item"])
        matched_refs = []

        for code in codes:
            if code in spec_index:
                matched_refs.append(spec_index[code]["source_id"])

        if matched_refs:
            # Deduplicate while preserving order
            seen = set()
            unique_refs = []
            for ref in matched_refs:
                if ref not in seen:
                    seen.add(ref)
                    unique_refs.append(ref)

            results.append({
                **row,
                "spec_reference": "; ".join(unique_refs),
                "match_type": "exact",
            })
        else:
            results.append({
                **row,
                "spec_reference": None,
                "match_type": "unmatched",
            })

    return results
