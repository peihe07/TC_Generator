#!/usr/bin/env python3
"""VS-CF-03 §六／§〇 —— 對 Revise2 之全簿收尾檢查與目標數字現算。

判準：
  §〇 七項目標數字（資料列、Requirement 覆蓋、未生成、PENDING、DR 引用、
       IMPL_GAP、tc_id 連續）
  §六.1 全形 `；`／`，`（作者側四欄）
  §六.2 `Set PROXI`
  §六.3 尾句號／行首尾空白／`$MESSAGE.Signal$`／方括號分類
  §六.4 Procedure ↔ ER 1:1、每列 ≥2 步
  §六.5 `specification_reference` 格式與升冪
  §六.6 Priority 值域、design_method 值域
  §六.7 §10.6 嚴格等價之重複比對（命中即回報，不自行合併）
  §六.8 row 13 之改寫（讀回自注訊號歸零）
另檢 R-S4：`test_item` 括號下半存在、無中文、同 Req 不逐字相同。
"""

from __future__ import annotations

import re
import sys
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
FEAT = ROOT / "features/vehicle_setting"
BOOK = FEAT / "sandbox/cfts044/cfts044_20260819_Revise2.xlsx"
SHEET = "Test Case Specification 測試用例規範"
SWRA = sorted((FEAT / "inputs").glob("*037*CFTS044*.xlsx"))

C = {"B": 2, "D": 4, "F": 6, "G": 7, "H": 8, "I": 9, "J": 10, "K": 11,
     "L": 12, "M": 13, "N": 14, "P": 16, "R": 18, "AH": 34}
AUTHOR = ("J", "K", "L", "M")
METHODS = {
    "功能測試 (Functional based ; no specific technique)",
    "狀態轉換 (State Transition Testing)", "決策表 (Decision Table Testing)",
    "等價劃分 (Equivalence Partitioning, EP)",
    "邊界值分析 (Boundary Value Analysis, BVA)",
    "組合測試 (Combinatorial Testing ; Pairwise / t-wise)",
    "情境 / 用例 (Scenario / Use Case Testing)", "負向測試 (Negative / Invalid)",
    "基礎故障注入 (Fault Injection Lite)",
}
WITHDRAWN_DR = ("DR-15", "DR-21", "DR-24", "DR-25", "DR-26", "DR-18", "DR-22")
PAREN_TAIL = re.compile(r"\([^)]{5,}\)\s*$", re.S)
CJK = re.compile(r"[　-〿一-鿿＀-￯]")

warnings.filterwarnings("ignore")


def lines(v) -> list[str]:
    return [x for x in str(v or "").split("\n") if x.strip()]


def main() -> int:
    ws = openpyxl.load_workbook(BOOK)[SHEET]
    rows = {r: {k: ws.cell(r, c).value for k, c in C.items()}
            for r in range(10, ws.max_row + 1)}
    data = [r for r, d in rows.items() if str(d["F"] or "").strip()]
    blob = lambda r: "\n".join(str(rows[r][k] or "")
                               for k in ("I", "J", "K", "L", "M", "AH"))
    fails: list[str] = []

    def chk(label: str, bad: list, extra: str = "") -> None:
        ok = not bad
        if not ok:
            fails.append(label)
        print(f"{label:38} {'PASS' if ok else 'FAIL'} "
              f"{('' if ok else str(bad[:12])) or extra}")

    # ------------------------------------------------------ §〇 目標數字
    leaves = set()
    for f in SWRA:
        w = openpyxl.load_workbook(f, read_only=True)["Analysis Report"]
        for r in range(8, w.max_row + 1):
            if w.cell(r, 1).value:
                leaves.add(str(w.cell(r, 1).value).strip())
    reqs = Counter(str(rows[r]["D"]).strip() for r in data)
    print("== §〇 目標數字 ==")
    print(f"{'資料列 243 → 241':38} {len(data)}")
    print(f"{'Requirement 覆蓋':38} {len(reqs)} / 237 "
          f"（皆在 037：{all(k in leaves for k in reqs)}）")
    notgen = [r for r in data if not str(rows[r]["I"] or "").strip()]
    print(f"{'未生成列 23 → 0':38} {len(notgen)} {notgen}")
    pend = [r for r in data if "PENDING" in blob(r)]
    print(f"{'PENDING 12 → 0':38} {len(pend)} {pend}")
    dr = [r for r in data if any(x in blob(r) for x in WITHDRAWN_DR)
          or "BLOCKED" in blob(r)]
    print(f"{'撤回 DR／BLOCKED 引用 → 0':38} {len(dr)} {dr}")
    gap = [r for r in data if "IMPL_GAP" in blob(r)]
    print(f"{'IMPL_GAP → 0':38} {len(gap)} {gap}")
    open_dr = sorted({m for r in data
                      for m in re.findall(r"dr_dependent = (DR-[\w-]+)", blob(r))})
    print(f"{'仍具名之未結 DR（非撤回清單）':38} {open_dr} "
          f"{len([r for r in data if 'dr_dependent' in blob(r)])} 列")
    ids = [str(rows[r]["F"]) for r in data]
    want = [f"NR1L-VehicleSetting-{i:03d}" for i in range(1, len(data) + 1)]
    print(f"{'tc_id 重賦連續':38} {'PASS' if ids == want else 'FAIL'}")
    print()

    print("== §六 全簿收尾 ==")
    chk("§〇 未生成 / PENDING / DR / IMPL_GAP",
        notgen + pend + dr + gap)
    chk("tc_id 連續", [] if ids == want else ["mismatch"])
    chk("No.# 連續",
        [r for i, r in enumerate(data) if rows[r]["B"] != i + 1])

    # §六.1
    chk("§六.1 全形 ；／，（作者側四欄）",
        [r for r in data
         if any("；" in str(rows[r][c] or "") or "，" in str(rows[r][c] or "")
                for c in AUTHOR)])
    # §六.2
    chk("§六.2 Set PROXI",
        [r for r in data if any("Set PROXI" in str(rows[r][c] or "")
                                for c in AUTHOR)])
    # §六.3
    chk("§六.3 行尾句號",
        [r for r in data if any(ln.strip().endswith((".", "。"))
                                for c in AUTHOR for ln in lines(rows[r][c]))])
    chk("§六.3 行首尾空白",
        [r for r in data
         if any(ln != ln.strip() for c in AUTHOR + ("I", "N", "AH")
                for ln in str(rows[r][c] or "").split("\n"))])
    chk("§六.3 $MESSAGE.Signal$ 三件組",
        [r for r in data if any(re.search(r"\$[A-Z0-9_]+\.[A-Za-z0-9_]+\$",
                                          str(rows[r][c] or ""))
                                for c in AUTHOR)])
    ui_bracket = [r for r in data
                  if any(re.search(r"(?:Press|Select|Check that the)\s+\[",
                                   str(rows[r][c] or "")) for c in AUTHOR)]
    chk("§六.3 方括號（UI 標籤式）", ui_bracket)
    val_bracket = sum(len(re.findall(r"\[[^\]]+\]", str(rows[r][c] or "")))
                      for r in data for c in AUTHOR)
    print(f"{'　└ §一(a) 值標籤 token（IN §11 例外）':38} {val_bracket} 處（不計違規）")

    # §六.4
    chk("§六.4 Procedure ↔ ER 1:1",
        [r for r in data
         if len(lines(rows[r]["L"])) != len(lines(rows[r]["M"]))])
    chk("§六.4 每列 ≥2 步", [r for r in data if len(lines(rows[r]["L"])) < 2])

    # §六.5
    bad_ref, unsorted_ref = [], []
    for r in data:
        v = [x.strip() for x in str(rows[r]["N"] or "").split("\n") if x.strip()]
        if not v:
            continue
        if not all(re.fullmatch(r"CFTS044-\d{7}", x) for x in v):
            bad_ref.append(r)
        if v != sorted(v):
            unsorted_ref.append(r)
    chk("§六.5 specification_reference 格式", bad_ref)
    chk("§六.5 specification_reference 升冪", unsorted_ref)
    empty_ref = [r for r in data if not str(rows[r]["N"] or "").strip()]
    print(f"{'　└ 無 CFTS044 錨之列（具名）':38} {empty_ref}")

    # §六.6
    chk("§六.6 Priority 值域",
        [r for r in data if str(rows[r]["P"]) not in ("P0", "P1", "P2", "P3")])
    chk("§六.6 design_method 值域",
        [r for r in data if str(rows[r]["R"]) not in METHODS])

    # §六.7 §10.6 嚴格等價
    sig = defaultdict(list)
    for r in data:
        key = tuple(str(rows[r][c] or "").strip() for c in ("I", "J", "K", "L", "M"))
        sig[key].append(r)
    dups = [v for v in sig.values() if len(v) > 1]
    # §六.7：命中即回報，不自行合併 —— 故列為 REPORT，不計入 lint 失敗
    print(f"{'§六.7 五欄嚴格等價重複':38} REPORT {dups}")

    # §六.8
    chk("§六.8 row 13 不讀回自注訊號",
        [] if "ESS_ENG_ST and check" not in str(rows[13]["L"] or "")
        else ["row 13"])

    # R-S4
    chk("R-S4 括號下半存在",
        [r for r in data if not PAREN_TAIL.search(str(rows[r]["I"] or "").strip())])
    chk("R-S4 括號下半無中文",
        [r for r in data
         if CJK.search(str(rows[r]["I"] or "").strip().split("\n")[-1])])
    by_req = defaultdict(list)
    for r in data:
        low = str(rows[r]["I"] or "").strip().split("\n")[-1]
        by_req[str(rows[r]["D"]).strip()].append((r, low))
    same = [v for v in by_req.values()
            if len(v) > 1 and len({x[1] for x in v}) < len(v)]
    chk("R-S4 同 Req 括號下半相異", same)

    print()
    print("總判：", "PASS" if not fails else f"FAIL（{len(fails)} 項）")
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
