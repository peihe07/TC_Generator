#!/usr/bin/env python3
"""VS-CF-01 §五-4 —— 對 Revise1 之 lint。

`vs_sl03_lint.py` 綁 VS-SL-03 之三簿與其 v3 dryrun 報告（列母體自報告取），
**無法直接對 CFTS044 跑**。本檔沿用其三項硬規之判準（逐字取其 regex），
列母體改以工作簿之資料列取，並加本包之 PENDING 計數 gate。

三項硬規：`test_item` 括號下半、行尾句號、設定項名之雙引號。
本包**只改 23 列**，故判定以「本包觸及之列」為準；全簿數字併陳供 F1／F2／F3 用。
"""

from __future__ import annotations

import re
import sys
import warnings
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import vs_cf01_apply as A  # noqa: E402

warnings.filterwarnings("ignore")

# —— 判準逐字取自 vs_sl03_lint.py ——
PAREN_TAIL = re.compile(r"\([^)]{5,}\)\s*$", re.S)
NAME = r"(?:[A-Z][A-Za-z0-9\-/&]*)(?:\s+(?:[A-Z0-9][A-Za-z0-9\-/&]*|with|and|or|for))*"
UNQUOTED = re.compile(rf"\bthe\s+(?!\")({NAME})\s+setting\b")
UNQUOTED_ER = re.compile(rf"^The\s+(?!\")({NAME})\s+setting\b")

COL_ITEM, COL_PRE, COL_PROC, COL_ER = 9, 10, 12, 13


def check(ws, rows: list[int]) -> dict:
    g = lambda r, c: ("" if ws.cell(r, c).value is None else str(ws.cell(r, c).value))
    no_paren, dot, noq, fullwidth = [], [], [], []
    for r in rows:
        item = g(r, COL_ITEM).strip()
        # NOT GENERATED 列全欄清空，不受三項硬規拘束
        if not item and not g(r, COL_PROC).strip():
            continue
        if not PAREN_TAIL.search(item):
            no_paren.append(r)
        for c in (COL_PRE, COL_PROC, COL_ER):
            if any(x.strip().endswith(".") for x in g(r, c).split("\n") if x.strip()):
                dot.append(r)
                break
        for c in (COL_PROC, COL_ER):
            if any(UNQUOTED.search(ln) or UNQUOTED_ER.search(ln)
                   for ln in g(r, c).split("\n")):
                noq.append(r)
                break
        if any("；" in g(r, c) for c in (COL_PROC, COL_ER)):
            fullwidth.append(r)
    return {"no_paren": no_paren, "dot": dot, "noq": noq, "fullwidth": fullwidth}


def main() -> int:
    ws = openpyxl.load_workbook(A.OUT)[A.SHEET]
    data_rows = [r for r in range(10, ws.max_row + 1)
                 if str(ws.cell(r, 6).value or "").strip()]

    touched = check(ws, A.TOUCHED)
    whole = check(ws, data_rows)

    print(f"Revise1：{A.OUT.name}")
    print(f"資料列 {len(data_rows)}；本包觸及 {len(A.TOUCHED)} 列\n")
    bad = 0
    for key, label in (("no_paren", "test_item 無括號下半"),
                       ("dot", "行尾句號"),
                       ("noq", "設定項未加雙引號"),
                       ("fullwidth", "全形分號")):
        t, w = touched[key], whole[key]
        bad += len(t)
        print(f"{label:16} 觸及列 {len(t):3} {'PASS' if not t else 'FAIL ' + str(t)}"
              f"   ｜全簿 {len(w):3}")

    pend = A.assert_gates(ws)["pending"]
    print(f"\nPENDING gate：25 → {len(pend)}  PASS")
    for x in pend:
        print("   ", x)
    print("\n總判：", "PASS" if bad == 0 else "FAIL")
    return 0 if bad == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
