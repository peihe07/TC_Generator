#!/usr/bin/env python3
"""VS-CF-01 §五 —— Revise1 之自檢與逐列逐欄 diff（上繳用）。

不改任何檔；只讀母本與 Revise1，輸出 diff 與 gate 結果。
"""

from __future__ import annotations

import sys
import warnings
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import vs_cf01_apply as A  # noqa: E402

warnings.filterwarnings("ignore")
INV = {v: k for k, v in A.C.items()}


def show(v) -> str:
    return "∅" if v is None or v == "" else repr(str(v))


def main() -> int:
    old = openpyxl.load_workbook(A.MASTER)[A.SHEET]
    new = openpyxl.load_workbook(A.OUT)[A.SHEET]
    total = 0
    for r in range(1, max(old.max_row, new.max_row) + 1):
        cells = []
        for c in range(1, max(old.max_column, new.max_column) + 1):
            a, b = old.cell(r, c).value, new.cell(r, c).value
            if a != b:
                cells.append((c, a, b))
        if not cells:
            continue
        total += len(cells)
        print("=" * 72)
        print(f"ROW {r}  {old.cell(r, A.C['F']).value}")
        for c, a, b in cells:
            print(f"  [{INV.get(c, c)}]")
            print(f"    - {show(a)}")
            print(f"    + {show(b)}")
    print("=" * 72)
    print(f"改動儲存格總數：{total}")
    print(f"母本列數 {old.max_row}／Revise1 列數 {new.max_row}（須相等）")
    assert old.max_row == new.max_row, "列數變動，違反本包「不刪列」"
    A.assert_gates(new)
    print("gate：L／M 步數相等、PENDING = 12 —— 全過")
    return 0


if __name__ == "__main__":
    sys.exit(main())
