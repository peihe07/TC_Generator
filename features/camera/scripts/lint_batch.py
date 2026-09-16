"""批次 json → lint 專用暫存簿 → `lint036.py --profile camera`。

**不寫回任何交付工作簿**（Phase 5 之寫回一律走 `backend/xlsx_surgical.py`，R-G3）。
本工具只在指定之暫存目錄組一本最小 FW036 簿供 lint 讀，欄序取母本第 9 列實測
（R-G48；Vehicle Model 七欄 T–Z 為 `Z` 檢查所必需）。

    python features/camera/scripts/lint_batch.py <暫存目錄> [批次目錄]

批次目錄預設 `features/camera/generated/pilot01`。
"""
import json, sys, warnings
from pathlib import Path
import openpyxl
warnings.filterwarnings("ignore")
sys.path.insert(0, "/Users/peihe/Work_Projects/TC_Generator/scripts")
import lint036 as L

ROOT = Path("/Users/peihe/Work_Projects/TC_Generator")
GEN = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "features/camera/generated/pilot01"
SP = Path(sys.argv[1])
OUT = SP / f"{GEN.name}_lint.xlsx"

HDR = {"B": "No.#", "C": "Requirement or Design ID (Polarion)",
       "D": "Requirement or Design ID", "E": "Test Case ID (TestRail)",
       "F": "Test Case ID", "G": "Test Group", "H": "Test Set", "I": "Test Item",
       "J": "Pre-Conditions", "K": "Input Test Data", "L": "Test procedure",
       "M": "Expected Result", "N": "Specification Reference ", "O": "Test Case Reference ID",
       "P": "Test Case Priority", "Q": "Estimated Test Time (mins)",
       "R": "Test Case Design Methods", "S": "Functional Safety",
       "T": "HDCC27\nAtl-Hi\n", "U": "DT27\nAtl-Hi\n", "V": "VF(ProMaster)637\nAtl-Mi",
       "W": "Commander (598)\nAtl-Mi", "X": "Regengade (5210)\nAtl-Mi",
       "Y": "Toro(2261)\nAtl-Mi", "Z": "Fastack (376)\nAtl-Mi",
       "AA": "Test Case Author\n測試案例作者", "AH": "Remarks\n備註"}
VM = {"T": "HDCC27", "U": "DT27", "V": "VF(ProMaster)637", "W": "Commander (598)",
      "X": "Regengade (5210)", "Y": "Toro(2261)", "Z": "Fastack (376)"}

wb = openpyxl.Workbook(); ws = wb.active
ws.title = "Test Case Specification 測試用例規範"
for col, txt in HDR.items():
    ws[f"{col}9"] = txt
row = 10
for p in sorted(GEN.glob("NR1L-RVC-*.json")):
    d = json.loads(p.read_text(encoding="utf-8")); t = d["tcs"][0]
    ws[f"B{row}"] = row - 9
    ws[f"D{row}"] = d["req_id"]; ws[f"F{row}"] = d["tc_id"]
    ws[f"G{row}"] = d["test_group"]; ws[f"H{row}"] = d["test_set"]
    ws[f"I{row}"] = t["test_item"]; ws[f"J{row}"] = t["pre_conditions"]
    ws[f"K{row}"] = t["input_test_data"]; ws[f"L{row}"] = t["test_procedure"]
    ws[f"M{row}"] = t["expected_result"]; ws[f"N{row}"] = t["specification_reference"]
    ws[f"O{row}"] = "NEW"; ws[f"P{row}"] = t["priority"]
    ws[f"R{row}"] = t["design_method"]; ws[f"S{row}"] = "NA"
    for col, name in VM.items():
        ws[f"{col}{row}"] = d["vehicle_model"][name]
    ws[f"AA{row}"] = "PeiPYHsu"
    row += 1
wb.save(OUT)

res = L.lint_workbook(OUT, profile="camera")
tot = 0
for r in res:
    print(f"sheet={r.sheet} data_rows={r.data_rows} violations={len(r.violations)}")
    for v in r.violations:
        tot += 1
        print(f"  [{v.check}] row {v.row} {v.tc_id} {v.field}: {v.detail} | {v.snippet[:70]}")
print(f"\n合計違規 {tot}")
print("counts:", {k: n for k, n in L.count_by_check(res, "camera").items() if n})
