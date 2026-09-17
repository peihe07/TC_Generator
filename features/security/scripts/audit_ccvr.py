#!/usr/bin/env python3
"""CCVR SYS2_Mapped 落地審計（SEC-09，唯讀）。

三張表：
  audit_conflicts.tsv —— `Mapping notes` 六衝突是否落 Remarks
  audit_scope.tsv     —— CP 負向 TC 之「service 層 ≠ reprogram」範圍說明
  audit_newtc.tsv     —— `New test cases` 47 條 vs sibling plan 查漏

不改任何 TC；產物只落 `features/security/data/audit_ccvr/`。
"""
from __future__ import annotations

import csv
import hashlib
import importlib.util
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "features" / "security" / "data"
OUT = DATA / "audit_ccvr"
CCVR = (ROOT / "sources" / "raw" / "ccvr_v27_sys2_mapped" /
        "21047_24_00081_Supplier_Cybesecurity_Component_Verification_Report-v2.7_SYS2_Mapped.xlsx")

_s = importlib.util.spec_from_file_location("g1", Path(__file__).parent / "gen_batch1.py")
g1 = importlib.util.module_from_spec(_s)
_s.loader.exec_module(g1)

# 審計對象：最新 merged 本（版本同生成側旗標）
import os                                                    # noqa: E402
MERGED = (ROOT / "features/security/sandbox/merged"
          / f'security_{os.environ.get("SEC_VER", "v05")}.xlsx')

# §1 六衝突（下放包 §1；C4 之涉及列由執行層依 notes 內容指定）
CONFLICTS = [
    ("C1", "Root and lifecycle conflicts（543 root 不可改 vs 226 TLS list 可更新 vs 546 identity-root update）",
     ["543", "226", "546"],
     ["SWE1-CertProvider-007", "SWE1-CertProvider-009"]),
    ("C2", "Rollback conflict（344 禁降級 vs Auth-Prog CS.93 case 7 允許同 rollback index 之舊版）",
     ["344"],
     ["SWE1-SRA-SECURITY-SWDL-004", "SWE1-SRA-SECURITY-SWDL-003"]),
    ("C3", "Cryptographic baseline（376/509/511 ≥112-bit、510 ≥128-bit 非對稱、514 RSA-2048）",
     ["376", "509", "510", "511", "514"],
     ["SWE1-KeyInsyall-009", "SWE1-KeyInsyall-010", "SWE1-LOGENC-004"]),
    ("C4", "517（單一 code-signing root）vs 520（per-market CAs）—— 同列於 notes row 12",
     ["517", "520"],
     ["SWE1-CertProvider-007", "SWE1-CertProvider-009"]),
    ("C5", "DCL 未落地（notes row 10：DCL not yet／PROD certs not yet／generic INTERNAL_ERROR 非充分證據）",
     ["529", "383", "532", "539"],
     ["SWE1-CertProvider-004", "SWE1-CertProvider-005"]),
    ("C6", "CSR 格式與 key-binding 證據缺（notes row 15）",
     ["562", "563", "564", "565"],
     ["SYSAD_SEC_ECUCERT_ECUCERT_SRV_EXPORTCSR_INTF"]),
]

# §2 判準：範圍說明之等義句式
SCOPE_PAT = re.compile(
    r"scope:\s*service-level|does not cover reprogramming rejection|"
    r"reprogram path owned by SWDL|service-level rejection", re.I)
# 越權斷言：ER 不得斷言「ECU 不被重刷」
OVERREACH = re.compile(
    r"reprogramming is rejected|ECU is not reprogrammed|is not reflashed|"
    r"no new executable becomes active|application is not installed", re.I)

# §3 軸類別（四類，下放包 §3 步驟 2）。判準透明化，供分析層複判：
#   「需求面」自 NEW 之 title／steps／expected 抽（緊式關鍵詞，避免 `each` 等泛詞誤觸）；
#   「覆蓋面」自該 SWE1 之 sibling 集合判 —— sibling 本身即變體拆分，故
#   variant 之覆蓋以「sibling ≥ 2」為準，negative 以「有負向 sibling／負向 Design Method」為準。
AXIS_KW = {
    "negative": r"reject|untrusted|invalid|wrong|tamper|spoof|disallow|unauthorized|"
                r"denied|blocked|revoke|cannot|not permitted|fails? before",
    "boundary": r"expired|not-yet-valid|validity period|boundary|rate limit|minimum|maximum|"
                r"exceed|\d+-character|timeout|within \d+",
    "variant": r"per-market|each active|each configured|each required|each role|each interface|"
               r"DEV and PROD|both required|variants|multi-|each of four",
    "environment": r"manufactur|build-time|build time|production state|after transition|"
                   r"reboot|power-cycle|power loss|secure boot|OTA|USB|dealer",
}
COV_NEG = r"reject|fail|invalid|untrusted|wrong|mismatch|not identical|no |denied|revoked|" \
          r"prevent|refuse|error"
COV_BND = r"boundary|range|order|validity|counter|limit|target|uniqueness"
COV_ENV = r"architecture|platform|media|path|environment|build|upgrade|reset|reboot|" \
          r"persistence|target|stage|location"


# SEC-09 review §二（SEC-10 落地）：落點 SWE1 指錯者之改判 —— 追加真正覆蓋該軸之列。
COVERED_BY_OTHER = {
    # 跨簽 bus image 之拒絕 = KI-003 之 non-platform（R1LRefresh）負向 sibling
    "NEW-BOOT-02": ["SWE1-KeyInsyall-003"],
    # SELinux 關閉失敗 = `NR1L-KI-022`（KI-012 sibling 1，avc denied）
    "NEW-OS-01": ["SWE1-KeyInsyall-012"],
}
# 已裁不拆者之處置（`ANOMALIES.md`）
DISPOSITION = {
    "NEW-CERT-02": "A-SEC-8 (037 無到期／稽核語意，不拆)",
    "NEW-ID-01": "A-SEC-9 (ECUCert 037 十三列無 reboot／recovery／restart／power cycle，不拆)",
    "NEW-NET-01": "A-26 (trace_matrix 指派已清除，改判 NO_SWE1)",
    # SEC-11 review §一 #2（SEC-12 落地）：SWDL-001 之 037 VC 無 rollback 負向要件，不拆。
    "NEW-UPD-01": "A-SEC-11 (037 SWDL-001 VC 無負向要件，不拆)",
}


def sha16(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()[:16]


def read_merged() -> list[dict]:
    ws = openpyxl.load_workbook(MERGED, data_only=True)[g1.SHEET]
    out, r = [], g1.FIRST_ROW
    while ws[f'{g1.COLS["tc_id"]}{r}'].value:
        out.append({k: (ws[f"{c}{r}"].value or "") for k, c in g1.COLS.items()})
        r += 1
    return out


def axes_of(text: str) -> set[str]:
    """需求面之軸（緊式關鍵詞）。"""
    import re as _re
    return {k for k, pat in AXIS_KW.items() if _re.search(pat, text, _re.I)}


def covered_axes(swes, sib, by_swe) -> set[str]:
    """覆蓋面之軸：以 sibling 之拆分事實判，不以字面比對。"""
    out: set[str] = set()
    n_sib = sum(len(sib.get(s, [])) for s in swes)
    tcs = [t for s in swes for t in by_swe.get(s, [])]
    if n_sib >= 2 or len(tcs) >= 2:
        out.add("variant")
    blob = " ".join(f'{p["axis"]} {p["lower_half(English)"]}'
                    for s in swes for p in sib.get(s, []))
    if re.search(COV_NEG, blob, re.I) or any("負向測試" in t["design"] for t in tcs):
        out.add("negative")
    if re.search(COV_BND, blob, re.I):
        out.add("boundary")
    if re.search(COV_ENV, blob, re.I):
        out.add("environment")
    return out


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    tcs = read_merged()
    by_swe = defaultdict(list)
    for t in tcs:
        by_swe[t["req_id"]].append(t)
    tm = {r["swe1_id"]: r for r in csv.DictReader(
        (DATA / "trace_matrix.tsv").open(encoding="utf-8"), delimiter="\t")}
    bo = {r["swe1_id"]: r for r in csv.DictReader(
        (DATA / "batch_order.tsv").open(encoding="utf-8"), delimiter="\t")}
    plan = list(csv.DictReader((DATA / "sibling_plan.tsv").open(encoding="utf-8"),
                               delimiter="\t"))
    plan += list(csv.DictReader((DATA / "sibling_plan_batch2.tsv").open(encoding="utf-8"),
                                delimiter="\t"))
    sib = defaultdict(list)
    for p in plan:
        sib[p["swe1_id"]].append(p)

    # ---- §1 ----
    with (OUT / "audit_conflicts.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["conflict", "sys2_ids", "expected_swe1", "tc_ids_found",
                    "remarks_has_note", "remarks_excerpt"])
        c1 = []
        for cid, _desc, ids, swes in CONFLICTS:
            found, hit, exc = [], "N", ""
            for s in swes:
                for t in by_swe.get(s, []):
                    found.append(t["tc_id"])
                    rm = t["remarks"]
                    if re.search(r"\bconflict:", rm) or any(
                            re.search(rf"SYS-RA-SEC-{i}\b|\b{i}\b", rm) for i in ids):
                        hit = "Y"
                        exc = exc or f'{t["tc_id"]}: {rm.splitlines()[0][:80]}'
            w.writerow([cid, ";".join(f"SYS-RA-SEC-{i}" for i in ids), ";".join(swes),
                        ";".join(found), hit, exc])
            c1.append((cid, hit, len(found)))

    # ---- §2 ----
    cp = [t for t in tcs if t["test_group"] == "Cert Provider"]
    neg = [t for t in cp if "負向測試" in t["design"]
           or t["req_id"] in ("SWE1-CertProvider-004", "SWE1-CertProvider-005")]
    over = []
    with (OUT / "audit_scope.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["tc_id", "swe1_id", "design_method", "remarks_has_scope_note",
                    "remarks_excerpt", "er_overreach_lines"])
        for t in sorted(neg, key=lambda x: x["tc_id"]):
            has = "Y" if SCOPE_PAT.search(t["remarks"]) else "N"
            ol = [ln.strip() for ln in t["er"].splitlines() if OVERREACH.search(ln)]
            over += [(t["tc_id"], ln) for ln in ol]
            w.writerow([t["tc_id"], t["req_id"], t["design"].split(" (")[0], has,
                        t["remarks"].splitlines()[0][:80], " ⏎ ".join(ol)])

    # ---- §3 ----
    ws = openpyxl.load_workbook(CCVR, read_only=True, data_only=True)["New test cases"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    # trace_matrix 之 new_test_ids（分析層 SEC-01 之指派）反查
    by_new = defaultdict(set)
    for sid, r in tm.items():
        for item in (r["new_test_ids"] or "").split(";"):
            for nid in item.split():
                if nid.startswith("NEW-"):
                    by_new[nid].add(sid)
    # SYS2 反查（獨立第二路）
    by_sys2 = defaultdict(set)
    for sid, r in tm.items():
        for i in (r["sys_ra_sec"] or "").split(";"):
            if i:
                by_sys2[i.strip()].add(sid)

    cls_count = defaultdict(int)
    gaps = []
    with (OUT / "audit_newtc.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["new_id", "title", "sys2", "swe1_ids", "class", "covering_tc_ids",
                    "gap_axis", "swe1_via_sys2_only", "required_axes", "covered_axes",
                    "disposition"])
        for r in rows:
            nid = (r[0] or "").strip()
            if not nid:
                continue
            title = (r[1] or "").strip()
            sys2 = [x.strip() for x in str(r[2] or "").split("\n") if x.strip()]
            via_new = set(by_new.get(nid, set()))
            via_sys2 = set().union(*(by_sys2.get(i, set()) for i in sys2)) if sys2 else set()
            # 指派之權威面為 `new_test_ids`（分析層 SEC-01 建立）；SYS2 反查只作交叉驗證
            # （SEC-09 §5-1：兩路一致，(b) 未增任何列）。A-26 清除後 `NEW-NET-01` 即無指派，
            # 其 SYS2 交集屬偶然，不得回填 → 改判 `NO_SWE1`。
            swes = set(via_new) | set(COVERED_BY_OTHER.get(nid, []))
            text = " ".join(str(x or "") for x in r[1:6])
            req_ax = axes_of(text)
            tc_ids = [t["tc_id"] for s in swes for t in by_swe.get(s, [])]
            cov_ax = covered_axes(swes, sib, by_swe)
            if not swes:
                cls = "NO_SWE1"
            elif all(bo.get(s, {}).get("disposition") == "D" for s in swes):
                cls = "DEFERRED"
            elif not tc_ids:
                cls = "DEFERRED"
            else:
                gap = req_ax - cov_ax
                # SEC-10 review §一 #2：軸未拆但**已裁不拆**且有 A-SEC 號者歸 `ACCEPTED_GAP`。
                cls = ("COVERED" if not gap else
                       "ACCEPTED_GAP" if nid in DISPOSITION else "AXIS_GAP")
            gap = sorted(req_ax - cov_ax) if swes and tc_ids else []
            if cls in ("AXIS_GAP", "ACCEPTED_GAP"):
                gaps.append((nid, cls, ";".join(gap)))
            cls_count[cls] += 1
            w.writerow([nid, title, ";".join(sys2), ";".join(sorted(swes)), cls,
                        ";".join(sorted(set(tc_ids))), ";".join(gap),
                        ";".join(sorted(via_sys2 - via_new)),
                        ";".join(sorted(req_ax)), ";".join(sorted(cov_ax)),
                        DISPOSITION.get(nid, "")])

    print(f"審計本：{MERGED.relative_to(ROOT)}  sha16 {sha16(MERGED)}")
    print(f"CCVR：{CCVR.name}  sha16 {sha16(CCVR)}")
    print("§1 六衝突:", {c: h for c, h, _ in c1}, "涉及 TC:",
          {c: n for c, _, n in c1})
    print(f"§2 CP 負向／CP-004-005 TC = {len(neg)}；scope 註 Y = "
          f"{sum(1 for t in neg if SCOPE_PAT.search(t['remarks']))}；越權斷言 = {len(over)}")
    for tc, ln in over:
        print("   越權:", tc, ln[:90])
    print(f"§3 NEW {sum(cls_count.values())} 條:", dict(cls_count))
    print("   軸未拆:", gaps)
    return 0


if __name__ == "__main__":
    sys.exit(main())
