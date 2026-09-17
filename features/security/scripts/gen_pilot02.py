#!/usr/bin/env python3
"""Pilot02 —— SEC-03 §5 之 12 TC（pilot01 全部重寫 ＋ 2 列）。

形制依 SEC-03 §1 量測（SWC 0708 ＋ pm_29 兩本一致）：
  Priority `P0`~`P3`／Design Method 為 `下拉選單` 之中英雙語全文／
  `Input Test Data` = `NA`／`Test Case Reference ID` = `NEW`／`Estimated Test Time` 空／
  ER 字面值 `"…"`（反引號與單引號兩本皆 0 次）。
內容依 R-SEC15(a)~(j)、R-SEC7、R-SEC6、R-SEC10。
"""
# superseded by gen_batch1.py (SEC-04); do not run — constants frozen per R-TM13
# （SEC-15 review §三：apk 檔名常數停在 R-SEC24 之前，僅加註不改值、不刪檔）
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from backend.xlsx_surgical import surgical_save          # noqa: E402

TEMPLATE = ROOT / "forms" / ("FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT "
                             "STLA Test Case Specification & Result_SWQT_20260817_ext.xlsx")
SHEET = "Test Case Specification 測試用例規範"
OUT_DIR = ROOT / "features" / "security" / "sandbox" / "pilot02"
OUT_XLSX = OUT_DIR / "security_pilot02.xlsx"
FIRST_ROW = 10

COLS = {"req_id": "D", "tc_id": "F", "test_group": "G", "test_set": "H",
        "test_item": "I", "pre": "J", "input": "K", "proc": "L", "er": "M",
        "spec": "N", "tc_ref": "O", "priority": "P", "design": "R",
        "fs": "S", "author": "AA", "remarks": "AH"}
VM_COLS = ["T", "U", "V", "W", "X", "Y", "Z"]
VM_VALUES = ["1", "1", "1", "0", "0", "1", "1"]
VM_NAMES = ["HDCC27", "DT27", "VF(ProMaster)637", "Commander (598)",
            "Regengade (5210)", "Toro(2261)", "Fastack (376)"]

TOKEN = {"Cert Provider": "SWE1-CertProvider-SWE1R1-V1.0",
         "Key Install": "SWE1-KeyInstall-SWE1R1-V1.0",
         "SAM": "SWE1-SAM-SWE1R1-V1.1",
         "ECU Cert": "SWE1_ECUCert_FM-WI-FSM-037-A03"}

# §1 量測所定之 Design Method 逐字值（＝ `下拉選單` sheet 之全文）
FUNC = "功能測試 (Functional based ; no specific technique)"
NEG = "負向測試 (Negative / Invalid)"
STATE = "狀態轉換 (State Transition Testing)"
PRIORITY = "P1"            # SWC 0708 多數值（190/286）

RUNNER = ("com.mitsubishielectric.ahu.efw.lib.melcocertprovider.libcertproviderservice.test"
          "/androidx.test.runner.AndroidJUnitRunner")
FQCN = ("com.mitsubishielectric.ahu.efw.lib.melcocertprovider.libcertproviderservice.test"
        ".CertProviderServiceManagerTest")
ADB_ROOT = "DUT is connected via ADB with root permission (Dev/Eng build)"
RUNNER_READY = "Test runner CertProviderServiceManagerTest.apk is installed"
ASSETS = "SecurityAssets/oem-certs/cert-provider/"
OPENSSL_VERIFY = ("$ openssl verify -verbose -CAfile RootCert.pem -untrusted L1.pem "
                  "-untrusted L2.pem -untrusted L3.pem SAMcert.pem")
ER_CHAIN_OK = 'The openssl command prints "SAMcert.pem: OK" on stdout'


def instrument(method: str) -> str:
    return f"$ adb shell am instrument -w -e class {FQCN}#{method} {RUNNER}"


def er_rc(method: str) -> str:
    return f'The instrumentation runner reports "OK (1 test)" for {method}'


TCS = [
    dict(tc_id="NR1L-CP-001", group="Cert Provider", tset="Chain Verification",
         req="SWE1-CertProvider-001", nrl="NRL-349384", design=FUNC,
         top="The Cert Provider feature shall verify the integrity and authenticity of an X.509 "
             "leaf certificate provided by an external application.",
         bot="(valid leaf certificate; chain verification passes)",
         pre=[ADB_ROOT, RUNNER_READY,
              f"A valid leaf certificate and its chain are available under {ASSETS}"],
         steps=[("Run the OpenSSL chain verification of the leaf certificate on the host",
                 OPENSSL_VERIFY, ER_CHAIN_OK),
                ("Run the Cert Provider instrumentation test for the FOTA MCPU certificate normal flow",
                 instrument("fotaMcpuCertTestNormalFlow"), er_rc("fotaMcpuCertTestNormalFlow")),
                ("Confirm the Cert Provider log after the instrumentation run",
                 "$ adb logcat -s MelcoCertProviderTest",
                 "The adb logcat -s MelcoCertProviderTest output contains no error entry "
                 "for the leaf certificate")],
         remarks=["source: 037 SWE1-CertProvider-001 Requirement Description and Verification Criteria; "
                  "Test_Items.txt #01",
                  "apk pairing: high (apk_pairing.tsv #01)",
                  "ER 3 is an absence assertion; the error-entry prefix is inferred from the "
                  "CS.212 row 1 sample and is not asserted verbatim"]),

    dict(tc_id="NR1L-CP-002", group="Cert Provider", tset="Chain Verification",
         req="SWE1-CertProvider-001", nrl="NRL-349384", design=NEG,
         top="The Cert Provider feature shall verify the integrity and authenticity of an X.509 "
             "leaf certificate provided by an external application.",
         bot="(broken certificate chain; verification is rejected)",
         pre=[ADB_ROOT, RUNNER_READY,
              f"A leaf certificate whose issuer is absent from the chain is available under {ASSETS}"],
         steps=[("Run the OpenSSL chain verification of the broken leaf certificate on the host",
                 OPENSSL_VERIFY,
                 'The openssl command prints "unable to get issuer certificate" on stdout'),
                ("Run the Cert Provider instrumentation test for the broken certificate flow",
                 instrument("fotaMcpuCertTestBrokenCert"), er_rc("fotaMcpuCertTestBrokenCert")),
                ("Confirm the Cert Provider log after the instrumentation run",
                 "$ adb logcat -s MelcoCertProviderTest",
                 "PENDING: X-h log keyword")],
         remarks=["source: 037 SWE1-CertProvider-001; Test_Items.txt #02",
                  "apk pairing: high (apk_pairing.tsv #02)",
                  "ER 3 is PENDING: the CS.212 row 1 string belongs to ecuCertTestBrokenChainCert "
                  "and must not be moved to the FOTA scenario (R-SEC15(a))"]),

    dict(tc_id="NR1L-CP-003", group="Cert Provider", tset="Field Matching",
         req="SWE1-CertProvider-002", nrl="NRL-349385", design=FUNC,
         top="The Cert Provider feature shall verify that the Subject field of the leaf certificate "
             "matches the expected value defined in the configuration or provided dynamically.",
         bot="(subject name identical to the configuration; certificate is accepted)",
         pre=[ADB_ROOT, RUNNER_READY, "SAM Dongle is inserted into the DUT",
              "Configuration files are available under " + ASSETS + "{TYPE}/cfgs/"],
         steps=[("Read the Subject field of the leaf certificate on the host",
                 "$ openssl x509 -in leaf.crt -text",
                 'The openssl x509 output shows a "Subject:" line identical to the subject name '
                 "in " + ASSETS + "{TYPE}/cfgs/"),
                ("Run the OpenSSL chain verification of the leaf certificate on the host",
                 OPENSSL_VERIFY, ER_CHAIN_OK),
                ("Run the Cert Provider instrumentation test for the SAM certificate normal flow",
                 instrument("samCertTestNormalFlow"), er_rc("samCertTestNormalFlow"))],
         remarks=["source: 037 SWE1-CertProvider-002 Requirement Description and Verification Criteria "
                  "(openssl x509 and cfgs path verbatim); Test_Items.txt #04",
                  "apk pairing: high (apk_pairing.tsv #04)",
                  "Code Signing cert profile is not cited: this row is a SAM-tree scenario "
                  "(SAM tree profile is reference only)"]),

    dict(tc_id="NR1L-CP-004", group="Cert Provider", tset="Field Matching",
         req="SWE1-CertProvider-002", nrl="NRL-349385", design=NEG,
         top="The Cert Provider feature shall verify that the Subject field of the leaf certificate "
             "matches the expected value defined in the configuration or provided dynamically.",
         bot="(subject name not identical to the configuration; certificate is rejected)",
         pre=[ADB_ROOT, RUNNER_READY, "SAM Dongle is inserted into the DUT",
              "Configuration files are available under " + ASSETS + "{TYPE}/cfgs/"],
         steps=[("Read the Subject field of the wrong-subject leaf certificate on the host",
                 "$ openssl x509 -in leaf.crt -text",
                 'The openssl x509 output shows a "Subject:" line that is not identical to the '
                 "subject name in " + ASSETS + "{TYPE}/cfgs/"),
                ("Run the OpenSSL chain verification of the wrong-subject leaf certificate on the host",
                 OPENSSL_VERIFY,
                 'The openssl command prints "SAMcert.pem: OK" on stdout, so only the Subject '
                 "field is not identical"),
                ("PENDING: X-e wrong-subject certificate + trigger", None,
                 "PENDING: X-h log keyword")],
         remarks=["source: 037 SWE1-CertProvider-002 Verification Criteria (correct/wrong subject branch)",
                  "Step 3 is PENDING: no DUT-side trigger exists for the wrong-subject variant; "
                  "host openssl is not a DUT trigger (R-SEC15(b))",
                  "channel_feasible neg branch = N (trace_summary 4b)"]),

    dict(tc_id="NR1L-CP-005", group="Cert Provider", tset="Revocation",
         req="SWE1-CertProvider-004", nrl="NRL-349387", design=NEG,
         top="For non-ECU verification scenarios (e.g., SAM, FOTA, Second Party APKs), the system "
             "shall perform revocation checks using Certificate Revocation Lists (CRL) preinstalled "
             "in the secure storage of the /odm partition.",
         bot="(certificate listed in the local revocation list; certificate is rejected)",
         pre=[ADB_ROOT, RUNNER_READY,
              'A revocation list "CRL.r0" is preinstalled in the /odm partition',
              f"A revoked leaf certificate is available under {ASSETS}"],
         steps=[("List the local revocation list in the /odm partition",
                 "$ adb shell ls -l /odm/etc/cert_store",
                 'The adb shell ls -l /odm/etc/cert_store output lists "CRL.r0"'),
                ("PENDING: X-e revoked certificate + trigger", None,
                 "PENDING: X-h log keyword")],
         remarks=["source: 037 SWE1-CertProvider-004 Requirement Description and Verification Criteria "
                  "(CRL and /odm partition verbatim)",
                  "DCL (SD.00015/03, DID 2031) is a separate mechanism; see CCVR Cert Val CS.98",
                  "Step 2 is PENDING: apk pairing #12 secondPartyLevelTwoCertTestNormalFlow is a "
                  "positive flow and does not trigger the revocation branch (R-SEC15(g))"]),

    dict(tc_id="NR1L-CP-006", group="Cert Provider", tset="Field Matching",
         req="SWE1-CertProvider-011", nrl="NRL-349394", design=FUNC,
         top="The Cert Provider feature shall support the validation of project-specific critical "
             "extensions within the leaf certificate.",
         bot="(extension value is the project identifier; certificate is accepted)",
         pre=[ADB_ROOT, RUNNER_READY,
              "A leaf certificate with the project-specific critical extension is available"],
         steps=[("Read the certificate extensions on the host",
                 "$ openssl x509 -in leaf.crt -text",
                 'The openssl x509 output shows "X509v3 Certificate Policies: critical" and a '
                 '"Policy:" line with the OID 1.3.6.1.4.1.57872.<…>'),
                ("Run the OpenSSL chain verification of the leaf certificate on the host",
                 OPENSSL_VERIFY, ER_CHAIN_OK),
                ("PENDING: X-e OID certificate + trigger", None,
                 "PENDING: X-h log keyword")],
         remarks=["source: 037 SWE1-CertProvider-011 Requirement Description and Verification Criteria; "
                  "CertProfile r36 (Code Signing tree, OID prefix verbatim)",
                  "OID suffix kept as a placeholder: CertProfile does not give the value",
                  "Step 3 is PENDING: apk_pairing.tsv has no method for this SWE1 row (R-SEC15(g))"]),

    dict(tc_id="NR1L-CP-007", group="Cert Provider", tset="Field Matching",
         req="SWE1-CertProvider-011", nrl="NRL-349394", design=NEG,
         top="The Cert Provider feature shall support the validation of project-specific critical "
             "extensions within the leaf certificate.",
         bot="(extension value is incorrect or the critical extension is missing; certificate is rejected)",
         pre=[ADB_ROOT, RUNNER_READY,
              "A leaf certificate without the project-specific critical extension is available"],
         steps=[("Read the certificate extensions of the wrong-OID certificate on the host",
                 "$ openssl x509 -in leaf.crt -text",
                 'The openssl x509 output shows no "X509v3 Certificate Policies: critical" line, '
                 'or a "Policy:" line whose OID is not 1.3.6.1.4.1.57872.<…>'),
                ("PENDING: X-e OID certificate + trigger", None,
                 "PENDING: X-h log keyword")],
         remarks=["source: 037 SWE1-CertProvider-011 Verification Criteria (second WHEN/THEN pair); "
                  "CertProfile r36",
                  "Step 2 is PENDING: no DUT-side trigger and no paired apk method (R-SEC15(b)(g))",
                  "channel_feasible neg branch = N (trace_summary 4b)"]),

    dict(tc_id="NR1L-ECUC-001", group="ECU Cert", tset="Certificate Lifecycle",
         req="SYSAD_SEC_ECUCERT_ECUCERT_SRV_EXPORTCSR_INTF", nrl="-", design=FUNC,
         top="Generate CSR message.",
         bot="(CSR is read over DID 2965 and exported to a file)",
         pre=[ADB_ROOT, "ADA Active Roles permit ADA+ protected data identifiers",
              "The ECU serial number has been written",
              "The CSR is generated at boot"],
         steps=[("Send UDS request DiagnosticSessionControl for the SystemSupplierSpecific session",
                 "$ 10 60", "Positive response is received: 50 60"),
                ("Send UDS request ReadDataByIdentifier for DID 2965 (CSR Read)",
                 "$ 22 29 65",
                 "Positive response is received: 62 29 65 <CSR bytes>, up to 1000 bytes"),
                ("Export the CSR file from the DUT to the host",
                 "$ adb pull /data/misc/ecuidentity/ecu.crt",
                 'The adb shell ls -l /data/misc/ecuidentity output lists "ecu.crt" on the DUT'),
                ("Check the exported CSR on the host",
                 "$ openssl req -in ecu.crt -noout -verify",
                 'The openssl command prints "verify OK" on stdout')],
         remarks=["source: R-SEC6(a) UDS step form; CS.00102 5.2.2.75 (SYS-RA-CS00102-415 to -418); "
                  "ECU Cert Test Steps PDF Step 2; CCVR ECU ID CS.165 item 12 (openssl req verbatim)",
                  "SWE ID pending; the SYSAD identifier is used instead",
                  "The consistency between the DID payload and ecu.crt is not asserted: "
                  "no comparison means is available in the delivered material"]),

    dict(tc_id="NR1L-KI-001", group="Key Install", tset="Install State",
         req="SWE1-KeyInsyall-011", nrl="-", design=STATE,
         top="The software shall provide a Binder-based interface via the Android Java Service to "
             "allow external components (specifically the DIAG service) to query the current key "
             "installation status.",
         bot="(installation completed; the status is INSTALLED (2))",
         pre=[ADB_ROOT, "Key installation has completed successfully on the DUT"],
         steps=[("Read the installation state file in the secure partition",
                 "$ adb shell od -t x1 /mnt/vendor/oemkeys/<installstate file>",
                 "The adb shell od -t x1 output shows the installstate value 2"),
                ("PENDING: X-f KeyInstall status test runner", None,
                 "PENDING: X-f KeyInstall status test runner")],
         remarks=["source: 037 SWE1-KeyInsyall-011 Requirement Description and Verification Criteria "
                  "1.1 and 3.1 (Binder getStatus verbatim)",
                  "RD ID spelling as delivered",
                  "The installstate file name is a placeholder: 037 states the secure partition "
                  "but not the file name. DID 22 FF 02 is a supplementary observation only and "
                  "is not a pass/fail criterion; the two-state versus four-state question is open"]),

    dict(tc_id="NR1L-KI-002", group="Key Install", tset="Install State",
         req="SWE1-KeyInsyall-011", nrl="-", design=NEG,
         top="The software shall provide a Binder-based interface via the Android Java Service to "
             "allow external components (specifically the DIAG service) to query the current key "
             "installation status.",
         bot="(unsuccessful installation; the status is ERROR (3))",
         pre=[ADB_ROOT, "An unsuccessful installation has been made on the DUT"],
         steps=[("Read the installation state file in the secure partition",
                 "$ adb shell od -t x1 /mnt/vendor/oemkeys/<installstate file>",
                 "The adb shell od -t x1 output shows the installstate value 3"),
                ("PENDING: X-f KeyInstall status test runner", None,
                 "PENDING: X-f KeyInstall status test runner")],
         remarks=["source: 037 SWE1-KeyInsyall-011 Verification Criteria 4.1 and 6.1 verbatim",
                  "RD ID spelling as delivered",
                  "Same placeholder and DID notes as the INSTALLED sibling"]),

    dict(tc_id="NR1L-SAM-001", group="SAM", tset="Error Handling",
         req="SWE1-SAM-0004", nrl="-", design=NEG,
         top="DebugAuth just logs an error, if verification or ID check failed.",
         bot="(AuthData verification fails; the installed AuthData is retained)",
         pre=[ADB_ROOT, "AuthData is already installed on the DUT",
              "SAM dongle is inserted into the DUT",
              "A USB drive containing an AuthData package with a mismatching signature is available"],
         steps=[("Record the installed AuthData before the verification attempt",
                 "$ adb shell ls -l /data/vendor/dauth",
                 "The adb shell ls -l /data/vendor/dauth output lists the installed AuthData"),
                ("Clear the log buffer before the verification attempt",
                 "$ adb logcat -c",
                 "The adb logcat -s logdog output is empty before the attempt"),
                ("Insert USB drive containing the invalid AuthData package into HU USB port",
                 None, "PENDING: X-h log keyword"),
                ("Record the installed AuthData after the verification attempt",
                 "$ adb shell ls -l /data/vendor/dauth",
                 "The adb shell ls -l /data/vendor/dauth output lists the same installed AuthData "
                 "as in step 1")],
         remarks=["source: 037 SWE1-SAM-0004 Requirement Description and Verification Criteria "
                  "3.1 and 3.2 (logdog assertion)",
                  "ER 3 is PENDING: the logdog keyword for this scenario is not given verbatim "
                  "in the delivered material",
                  "channel_feasible = N (trace_summary 4b)"]),

    dict(tc_id="NR1L-SAM-002", group="SAM", tset="AuthData Verification",
         req="SWE1-SAM-0007", nrl="-", design=FUNC,
         top="DebugAuth shall check SAM certificate as part of AuthData verification.",
         bot="(SAM certificate is x509 PEM, anchors to the OEM issued root, and carries the SAM subject)",
         pre=[ADB_ROOT, RUNNER_READY, "SAM dongle is inserted into the DUT",
              "An AuthData package is available on a USB drive"],
         steps=[("Read the SAM certificate format and subject on the host",
                 "$ openssl x509 -in SAMcert.pem -text",
                 "The openssl x509 output shows the certificate in x509 PEM format and a "
                 '"Subject:" line intended for SAM and specific market'),
                ("Run the OpenSSL chain verification of the SAM certificate against the OEM issued "
                 "root certificate", OPENSSL_VERIFY, ER_CHAIN_OK),
                ("Run the Cert Provider instrumentation test for the SAM certificate normal flow",
                 instrument("samCertTestNormalFlow"), er_rc("samCertTestNormalFlow"))],
         remarks=["source: 037 SWE1-SAM-0007 Verification Criteria 3.1, 3.2 and 3.3 verbatim; "
                  "Test_Items.txt #04",
                  "apk pairing: high (apk_pairing.tsv #04)",
                  "batch: manual override per R-SEC12. SAM tree cert profile is reference only"]),
]


def render(tc: dict) -> dict:
    proc, er = [], []
    for n, (desc, cmd, e) in enumerate(tc["steps"], 1):
        proc.append(f"{n}. {desc}")
        if cmd:
            proc.append(cmd)
        er.append(f"{n}. {e}")
    remarks = list(tc["remarks"])
    if tc["nrl"] != "-":
        remarks.insert(1, f'Polarion: {tc["nrl"]}')
    return {"req_id": tc["req"], "tc_id": tc["tc_id"], "test_group": tc["group"],
            "test_set": tc["tset"], "test_item": f'{tc["top"]}\n{tc["bot"]}',
            "pre": "\n".join(f"{i}. {x}" for i, x in enumerate(tc["pre"], 1)),
            "input": "NA", "proc": "\n".join(proc), "er": "\n".join(er),
            "spec": f'{TOKEN[tc["group"]]}_{tc["req"]}', "tc_ref": "NEW",
            "priority": PRIORITY, "design": tc["design"], "fs": "No",
            "author": "PeiPYHsu", "remarks": "\n".join(remarks)}


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    for n, tc in enumerate(TCS):
        row, r = FIRST_ROW + n, render(tc)
        for key, col in COLS.items():
            ws[f"{col}{row}"] = r[key]
        for col, val in zip(VM_COLS, VM_VALUES):
            ws[f"{col}{row}"] = val
        (OUT_DIR / f'{tc["tc_id"]}.json').write_text(
            json.dumps({**r, "vehicle_model": dict(zip(VM_NAMES, VM_VALUES))},
                       ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = surgical_save(wb, TEMPLATE, OUT_XLSX)
    print(f"pilot02：{len(TCS)} TC → {OUT_XLSX.relative_to(ROOT)}")
    print("  surgical_save:", {k: v for k, v in report.items() if k != "members_patched"})
    return 0


if __name__ == "__main__":
    sys.exit(main())
