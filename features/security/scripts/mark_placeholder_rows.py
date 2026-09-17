#!/usr/bin/env python3
"""佔位列整列淡橘底色（SEC-17；R-SEC25 之六本交付形制）。

**只改填色，不改任何文字**。做法與 `backend.xlsx_surgical.surgical_restyle` 同宗：
不改既有 `<xf>`（那會連帶重塗所有共用該 id 之格），而是
  1. 於 `styles.xml` 之 `<fills>` 追加一個 solid fill；
  2. 為每個「來源 style id」衍生一個帶該 fill 之新 `<xf>`，追加於 `<cellXfs>`；
  3. 只把目標列之 A~<表頭末欄> 重新指向新 id（缺格者補 `<c r=… s=…/>`）。
其餘 zip 成員逐位元組複製 —— `sharedStrings.xml` 不動即文字不可能變（另以逐格比對複驗）。

下放包 §2 令「`surgical_save` 之後以 openpyxl 逐格 PatternFill 寫回、再存」：
openpyxl 之 `save()` 會重建整包（R16-1：x14 DV 全失、zip 成員增減），
與 §0／§5 之「DV／zip 成員數不得變」直接衝突，故改用本法（見上繳包 §5）。
"""
from __future__ import annotations

import csv
import importlib.util
import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from backend.xlsx_surgical import (CELLXFS_RE, ROW_RE, STYLES_MEMBER,  # noqa: E402
                                   XF_RE, StructureError, col_to_idx,
                                   idx_to_col, sheet_members, verify_structure)

_s = importlib.util.spec_from_file_location("g1", Path(__file__).parent / "gen_batch1.py")
g1 = importlib.util.module_from_spec(_s)
_s.loader.exec_module(g1)

DATA = ROOT / "features" / "security" / "data"
SB = ROOT / "features" / "security" / "sandbox"
DELIVERY = SB / "delivery"
# §1：淡橘 `FCE4D6`（Excel「橙色, 輔色 2, 較淺 80%」）。量測見上繳包 §1。
FILL_RGB = "FFFCE4D6"
RE_PLACEHOLDER = re.compile(r"<[^<>]+provided by [^<>]+\(X-[a-z0-9-]+\)>")
FILLS_RE = re.compile(r'(<fills count=")(\d+)(">)(.*?)(</fills>)', re.S)
CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)


def target_rows(book: Path) -> tuple[list[int], dict[int, str], int]:
    """含佔位之資料列（1-based）、列 → TC ID、表頭最末欄索引。"""
    ws = openpyxl.load_workbook(book, data_only=True)[g1.SHEET]
    hdr = g1.FIRST_ROW - 1
    last_col = max(i for i in range(1, ws.max_column + 1) if ws.cell(hdr, i).value)
    rows, ids, r = [], {}, g1.FIRST_ROW
    while ws[f'{g1.COLS["tc_id"]}{r}'].value:
        blob = "\n".join(str(ws[f'{g1.COLS[k]}{r}'].value or "")
                         for k in ("pre", "input", "proc", "er"))
        if RE_PLACEHOLDER.search(blob):
            rows.append(r)
            ids[r] = ws[f'{g1.COLS["tc_id"]}{r}'].value
        r += 1
    return rows, ids, last_col


def _append_fill(styles: str) -> tuple[str, int]:
    m = FILLS_RE.search(styles)
    if m is None:
        raise StructureError("styles.xml 無 <fills> 區塊")
    count = int(m.group(2))
    new = (f'<fill><patternFill patternType="solid"><fgColor rgb="{FILL_RGB}"/>'
           '<bgColor indexed="64"/></patternFill></fill>')
    body = m.group(4) + new
    return (styles[:m.start()] + m.group(1) + str(count + 1) + m.group(3) + body
            + m.group(5) + styles[m.end():], count)


def _derive_xfs(styles: str, src_ids: set[int], fill_id: int) -> tuple[str, dict[int, int]]:
    m = CELLXFS_RE.search(styles)
    if m is None:
        raise StructureError("styles.xml 無 <cellXfs> 區塊")
    xfs = XF_RE.findall(m.group(4))
    if len(xfs) != int(m.group(2)):
        raise StructureError(f"cellXfs count={m.group(2)} 但解析得 {len(xfs)} 個 <xf>")
    mapping, appended = {}, []
    for sid in sorted(src_ids):
        xf = xfs[sid] if sid < len(xfs) else '<xf xfId="0"/>'
        xf = re.sub(r'\sfillId="\d+"', "", xf)
        xf = re.sub(r'\sapplyFill="[01]"', "", xf)
        xf = xf.replace("<xf ", f'<xf fillId="{fill_id}" applyFill="1" ', 1)
        mapping[sid] = len(xfs) + len(appended)
        appended.append(xf)
    body = m.group(4) + "".join(appended)
    total = len(xfs) + len(appended)
    return (styles[:m.start()] + m.group(1) + str(total) + m.group(3) + body
            + m.group(5) + styles[m.end():], mapping)


def _row_style_ids(xml: str, rows: set[int], last_col: int) -> set[int]:
    out: set[int] = set()
    for m in ROW_RE.finditer(xml):
        if int(m.group(1)) not in rows:
            continue
        for c in CELL_RE.finditer(m.group(3) or ""):
            if col_to_idx(c.group(1)) <= last_col:
                s = re.search(r'\ss="(\d+)"', c.group(3) or "")
                out.add(int(s.group(1)) if s else 0)
        rs = re.search(r'\ss="(\d+)"', m.group(2) or "")
        out.add(int(rs.group(1)) if rs else 0)
    return out


def _paint_rows(xml: str, rows: set[int], last_col: int,
                mapping: dict[int, int]) -> tuple[str, int]:
    painted = 0

    def row_sub(m: re.Match) -> str:
        nonlocal painted
        rno = int(m.group(1))
        if rno not in rows:
            return m.group(0)
        attrs, body = m.group(2) or "", m.group(3) or ""
        cells: dict[int, str] = {}
        for c in CELL_RE.finditer(body):
            cells[col_to_idx(c.group(1))] = c.group(0)
        rs = re.search(r'\ss="(\d+)"', attrs)
        row_style = int(rs.group(1)) if rs else 0
        for idx in range(1, last_col + 1):
            coord = f"{idx_to_col(idx)}{rno}"
            cell = cells.get(idx)
            if cell is None:
                cells[idx] = f'<c r="{coord}" s="{mapping[row_style]}"/>'
                painted += 1
                continue
            s = re.search(r'\ss="(\d+)"', cell)
            sid = int(s.group(1)) if s else 0
            new = mapping[sid]
            cells[idx] = (re.sub(r'\ss="\d+"', f' s="{new}"', cell, count=1) if s
                          else cell.replace(f'<c r="{coord}"', f'<c r="{coord}" s="{new}"', 1))
            painted += 1
        new_body = "".join(cells[k] for k in sorted(cells))
        return f'<row r="{rno}"{attrs}>{new_body}</row>'

    return ROW_RE.sub(row_sub, xml), painted


def mark(src: Path, out: Path) -> dict:
    rows, ids, last_col = target_rows(src)
    members = sheet_members(src)
    member = members[g1.SHEET]
    with zipfile.ZipFile(src) as z:
        styles = z.read(STYLES_MEMBER).decode("utf-8")
        sheet = z.read(member).decode("utf-8")
    styles, fill_id = _append_fill(styles)
    src_ids = _row_style_ids(sheet, set(rows), last_col)
    styles, mapping = _derive_xfs(styles, src_ids, fill_id)
    sheet, painted = _paint_rows(sheet, set(rows), last_col, mapping)

    out.parent.mkdir(parents=True, exist_ok=True)
    patched = {STYLES_MEMBER: styles, member: sheet}
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = (patched[info.filename].encode("utf-8")
                    if info.filename in patched else zin.read(info.filename))
            zout.writestr(info, data)
    report = verify_structure(src, out, set(patched))
    # 文字面複驗：逐格比對來源與輸出
    a = openpyxl.load_workbook(src, data_only=True)[g1.SHEET]
    b = openpyxl.load_workbook(out, data_only=True)[g1.SHEET]
    text_diff = sum(1 for r in range(1, a.max_row + 1) for c in range(1, a.max_column + 1)
                    if a.cell(r, c).value != b.cell(r, c).value)
    report.update({"rows": rows, "ids": ids, "painted_cells": painted,
                   "last_col": idx_to_col(last_col), "text_diff": text_diff})
    return report


def main() -> int:
    books = sorted(DELIVERY.glob("*Security_*_20260917.xlsx"))
    assert len(books) == 6, f"交付本 {len(books)} 本"
    sup = DELIVERY / "superseded"
    sup.mkdir(exist_ok=True)
    rows_out: list[tuple[str, str, int, str]] = []
    ph = {}
    for r in csv.DictReader((DATA / "placeholder_summary_v09.tsv").open(encoding="utf-8"),
                            delimiter="\t"):
        ph.setdefault((r["workbook"], r["tc_id"]), set()).add(r["x_token"])
    total = 0
    for book in books:
        comp = book.name.split("Security_")[1].split("_20260917")[0]
        tmp = book.with_name(book.stem + ".v11.xlsx")
        rep = mark(book, tmp)
        if rep["text_diff"]:
            raise StructureError(f"{comp}: 文字 diff {rep['text_diff']} ≠ 0")
        shutil.move(str(book), str(sup / f"{book.stem}_v10{book.suffix}"))
        shutil.move(str(tmp), str(book))
        total += len(rep["rows"])
        for r in rep["rows"]:
            tid = rep["ids"][r]
            rows_out.append((comp, tid, r, ";".join(sorted(ph.get((comp, tid), {"?"})))))
        print(f"  {comp:14s} 塗色列 {len(rep['rows']):2d}／格 {rep['painted_cells']:4d}"
              f"（A~{rep['last_col']}）｜文字 diff {rep['text_diff']}"
              f"｜members {rep['members']}｜dv {rep['dv_counts'][sheet_members(book)[g1.SHEET]]}")
    all_src = SB / "merged" / "security_v10_all.xlsx"
    all_out = SB / "merged" / "security_v11_all.xlsx"
    rep = mark(all_src, all_out)
    print(f"  {'_all':14s} 塗色列 {len(rep['rows']):2d}／格 {rep['painted_cells']:4d}"
          f"｜文字 diff {rep['text_diff']}")
    with (DATA / "placeholder_rows.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["workbook", "tc_id", "row", "tokens"])
        for row in rows_out:
            w.writerow(row)
    shutil.copy(DATA / "placeholder_rows.tsv", DELIVERY / "placeholder_rows.tsv")
    print(f"  合計塗色 TC {total}；placeholder_rows.tsv {len(rows_out)} 列")
    return 0


if __name__ == "__main__":
    sys.exit(main())
