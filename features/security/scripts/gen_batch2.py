#!/usr/bin/env python3
"""batch 2 產出（SEC-05 §4）—— 產出群 13 列 → sibling → TC。冪等。

`sibling_plan_batch2.tsv` 於本檔一併產出（§4.2 先出計畫再寫 TC）。
CAN 步驟依 **R-SEC18**：raw／label 逐字取 `forms/` 之 DBC `VAL_`；查無者 `<…>` 佔位。
"""
from __future__ import annotations

import csv
import importlib.util
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from backend.xlsx_surgical import surgical_save          # noqa: E402

_s = importlib.util.spec_from_file_location("g1", Path(__file__).parent / "gen_batch1.py")
g1 = importlib.util.module_from_spec(_s)
_s.loader.exec_module(g1)

DATA = ROOT / "features" / "security" / "data"
OUT_DIR = ROOT / "features" / "security" / "sandbox" / "batch2"
OUT_XLSX = OUT_DIR / f"security_batch2_{g1.VER}.xlsx"

VM_HI = ["1", "1", "0", "0", "0", "0", "0"]      # HDCC27 / DT27
VM_MI = ["0", "0", "1", "0", "0", "1", "1"]      # VF637 / Toro / Fastback
BVA = "邊界值分析 (Boundary Value Analysis, BVA)"
DAUTH = "/data/vendor/dauth"

# ---- R-SEC18(a)(b)：DBC 逐字值（執行層自 forms/ 之 DBC 取，出處記於 Remarks）
CAN = {
    "hi": ("BCM_FD_10", "CmdIgnSts", "4", "RUN", "PDT27_E2A_R1_FDCAN8.dbc", "0x481"),
    "mi": ("STATUS_BH_BCM2", "CmdIgnSts", "4", "RUN", "P363_BH-CAN [07338]_3A_R2.dbc", "0x46C"),
}
# R-SEC18 amend(a)：`POWER_MODE_STS ↔ PowerModeSts` 之對照已刪（A-19：分析層推測橋，
# DBC 證實 `PowerModeSts` 之 `VAL_` 無 resume 值）。resume 分支改走 R-SEC19 之 PENDING。


def send_can(ee: str):
    msg, sig, raw, label, dbc, cid = CAN[ee]
    return (f"Send the ignition signal on the {'Atl-Hi' if ee == 'hi' else 'Atl-Mi'} bus",
            f"Send CAN: {msg}.{sig} = {raw} ({label})",
            f"{msg}.{sig} = {raw} ({label}) is sent")


def cat(path, er):
    return (f"Read {path} on the DUT", f"$ adb shell cat {path}", er)


ls, od, logcat, pend, phys = g1.ls, g1.od, g1.logcat, g1.pend, g1.phys
ssl_verify, apk = g1.ssl_verify, g1.apk
XG = "X-g SAM AuthData"

# (swe1, sibling_no, lower, axis, priority, design, steps, vm)
SPEC = [
 ("SWE1-SAM-0003", 1, "AuthData at the root of the USB external memory is received", "media layout",
  "P1", g1.FUNC, [phys("Insert USB drive containing the AuthData package into HU USB port",
                       f"PENDING: {XG} reception evidence"),
                  ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists the received AuthData")], None),
 ("SWE1-SAM-0003", 2, "AuthData not at the root of the USB external memory is not received", "media layout",
  "P1", g1.NEG, [pend(f"{XG} package outside the USB root + trigger"),
                 ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists no newly received AuthData")], None),
 ("SWE1-SAM-0005", 1, "manifest, SAM certificate and manifest signature are all valid", "verification outcome",
  "P0", g1.FUNC, [("Read the SAM certificate on the host", "$ openssl x509 -in SAMcert.pem -text",
                   'The openssl x509 output shows the certificate in x509 PEM format'),
                  ssl_verify(True),
                  pend(f"{XG} package that passes verification + trigger"),
                  logcat("dauth", "The adb logcat -s dauth output contains the successful verification entry")], None),
 ("SWE1-SAM-0005", 2, "manifest signature is invalid; verification fails", "verification outcome",
  "P0", g1.NEG, [pend(f"{XG} package with an invalid manifest signature + trigger"),
                 logcat("dauth", "PENDING: X-h log keyword")], None),
 ("SWE1-SAM-0006", 1, "head unit serial number is the same as the ECUId parameter", "serial match",
  "P0", g1.FUNC, [cat("/sys/devices/soc0/serial_number",
                      "The adb shell cat /sys/devices/soc0/serial_number output shows the head unit serial number"),
                  pend(f"{XG} package issued for this serial number + trigger"),
                  logcat("dauth", "The adb logcat -s dauth output contains the serial number match entry")], None),
 ("SWE1-SAM-0006", 2, "head unit serial number differs from the ECUId parameter", "serial match",
  "P0", g1.NEG, [pend(f"{XG} package issued for another serial number + trigger"),
                 logcat("dauth", "PENDING: X-h log keyword")], None),
 ("SWE1-SAM-0008", 1, "manifest, signature and certificate files follow the SAM package naming and format",
  "package format", "P2", g1.FUNC,
  [ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists the manifest, the signature and the certificate files"),
   ("Read the SAM certificate format on the host", "$ openssl x509 -in SAMcert.pem -text",
    "The openssl x509 output shows the certificate in x509 PEM format")], None),
 ("SWE1-SAM-0008", 2, "manifest file format is invalid; the package is not accepted", "package format",
  "P2", g1.NEG, [pend(f"{XG} package with a malformed manifest + trigger"),
                 logcat("dauth", "PENDING: X-h log keyword")], None),
 ("SWE1-SAM-0009", 1, "Install is true and SAMType is not Logging; AuthData is copied to the secure storage",
  "install option", "P1", g1.FUNC,
  [pend(f"{XG} package with Install true and a non-Logging SAMType + trigger"),
   ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists the copied AuthData")], None),
 ("SWE1-SAM-0009", 2, "SAMType is Logging; AuthData is not copied to the secure storage", "install option",
  "P1", g1.NEG, [pend(f"{XG} package with a Logging SAMType + trigger"),
                 ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists no copied AuthData")], None),
 ("SWE1-SAM-0010", 1, "TimeStamp is greater than or equal to the stored value; AuthData is accepted",
  "timestamp order", "P0", g1.FUNC,
  [cat(f"{DAUTH}/ts", f"The adb shell cat {DAUTH}/ts output shows the stored timestamp"),
   pend(f"{XG} package with a newer TimeStamp + trigger"),
   cat(f"{DAUTH}/ts", f"The adb shell cat {DAUTH}/ts output shows the updated timestamp")], None),
 ("SWE1-SAM-0010", 2, "TimeStamp is older than the stored value; AuthData is rejected", "timestamp order",
  "P0", g1.NEG, [cat(f"{DAUTH}/ts", f"The adb shell cat {DAUTH}/ts output shows the stored timestamp"),
                 pend(f"{XG} package with an older TimeStamp + trigger"),
                 cat(f"{DAUTH}/ts", f"The adb shell cat {DAUTH}/ts output shows the unchanged timestamp")], None),
 ("SWE1-SAM-0011", 1, "verification of the stored AuthData fails; the stored AuthData is deleted",
  "failure handling", "P0", g1.NEG,
  [ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists the stored AuthData", "Record"),
   pend(f"{XG} stored package that fails verification + trigger"),
   ls(DAUTH, f"The adb shell ls -l {DAUTH} output lists no stored AuthData")], None),
 ("SWE1-SAM-0012", 1, "number of ignition cycles is 0 and does not exceed ValidityCounter", "ValidityCounter boundary",
  "P0", BVA, [cat(f"{DAUTH}/vc", f"The adb shell cat {DAUTH}/vc output shows the ValidityCounter value 0"),
              pend(f"{XG} ignition cycle within the ValidityCounter + trigger"),
              logcat("dauth", "The adb logcat -s dauth output contains the AuthData in-use entry")], None),
 ("SWE1-SAM-0012", 2, "number of ignition cycles is 1 and does not exceed ValidityCounter", "ValidityCounter boundary",
  "P0", BVA, [cat(f"{DAUTH}/vc", f"The adb shell cat {DAUTH}/vc output shows the ValidityCounter value 1"),
              pend(f"{XG} ignition cycle within the ValidityCounter + trigger"),
              logcat("dauth", "The adb logcat -s dauth output contains the AuthData in-use entry")], None),
 ("SWE1-SAM-0012", 3, "number of ignition cycles reaches 65535 and exceeds ValidityCounter",
  "ValidityCounter boundary", "P0", BVA,
  [cat(f"{DAUTH}/vc", f"The adb shell cat {DAUTH}/vc output shows the ValidityCounter value 65535"),
   logcat("dauth", "PENDING: X-h log keyword")], None),
 ("SWE1-SAM-0013", 1, "head unit boots up on Atl-Hi; disabling is notified to all target functions",
  "EE architecture", "P1", g1.STATE, [send_can("hi"),
   logcat("dauth", "The adb logcat -s dauth output contains the OFF notification to the target functions")], "hi"),
 ("SWE1-SAM-0013", 2, "head unit boots up on Atl-Mi; disabling is notified to all target functions",
  "EE architecture", "P1", g1.STATE, [send_can("mi"),
   logcat("dauth", "The adb logcat -s dauth output contains the OFF notification to the target functions")], "mi"),
 ("SWE1-SAM-0013", 3, "head unit resumes from suspend on Atl-Hi; disabling is notified", "EE architecture",
  "P1", g1.STATE, [pend("X-n suspend/resume trigger method"),
   logcat("dauth", "The adb logcat -s dauth output contains the OFF notification to the target functions")], "hi"),
 ("SWE1-SAM-0013", 4, "head unit resumes from suspend on Atl-Mi; disabling is notified", "EE architecture",
  "P1", g1.STATE, [pend("X-n suspend/resume trigger method"),
   logcat("dauth", "The adb logcat -s dauth output contains the OFF notification to the target functions")], "mi"),
 ("SWE1-SAM-0014", 1, "verification succeeds; the target functions are notified according to SAMType",
  "notification trigger", "P1", g1.FUNC,
  [pend(f"{XG} package that passes verification + trigger"),
   logcat("dauth", "The adb logcat -s dauth output contains the ON notification to the target functions")], None),
 ("SWE1-SAM-0016", 1, "target function connects while the status is ON; the current status is notified",
  "current status", "P1", g1.FUNC,
  [pend(f"{XG} target function connection + trigger"),
   logcat("dauth", "The adb logcat -s dauth output contains the current status notification")], None),
 ("SWE1-SAM-0016", 2, "target function connects while the status is OFF; the current status is notified",
  "current status", "P1", g1.FUNC,
  [pend(f"{XG} target function connection + trigger"),
   logcat("dauth", "The adb logcat -s dauth output contains the current status notification")], None),
 ("SWE1-SAM-0017", 1, "cold boot on Atl-Hi; the sequence ID is set to 0", "EE architecture",
  "P1", g1.STATE, [send_can("hi"),
   cat(f"{DAUTH}/seqId", f"The adb shell cat {DAUTH}/seqId output shows the sequence ID 0")], "hi"),
 ("SWE1-SAM-0017", 2, "cold boot on Atl-Mi; the sequence ID is set to 0", "EE architecture",
  "P1", g1.STATE, [send_can("mi"),
   cat(f"{DAUTH}/seqId", f"The adb shell cat {DAUTH}/seqId output shows the sequence ID 0")], "mi"),
 ("SWE1-LOGENC-006", 1, "encryption triggered by the getsslog command produces the encrypted pair",
  "trigger path", "P2", g1.FUNC,
  [("Trigger the log snapshot on the DUT", "$ adb shell getsslog",
    "The adb shell getsslog command completes and returns to the prompt"),
   ls("/data/vendor/logdog",
      "The adb shell ls -l /data/vendor/logdog output lists the encrypted log and the encrypted key")], None),
 ("SWE1-LOGENC-006", 2, "encryption triggered by the special files on the USB drive produces the encrypted pair",
  "trigger path", "P2", g1.FUNC,
  [ls("/data/vendor/logdog",
      "The adb shell ls -l /data/vendor/logdog output lists the encrypted log and the encrypted key")], None),
 ("SWE1-LOGENC-006", 3, "symmetric key differs between two encryption iterations", "key uniqueness",
  "P2", g1.FUNC,
  [("Decrypt the first encrypted pair on the host",
    "$ ./logdecrypt_Ver2.sh HUssss_YYYYMMDD_HHMMSS_XXXX.enc HUssss_YYYYMMDD_HHMMSS_XXXX_enckey.enc",
    "The logdecrypt_Ver2.sh tool creates the folder HUssss_YYYYMMDD_HHMMSS_XXXX on the host"),
   ("Decrypt the second encrypted pair on the host",
    "$ ./logdecrypt_Ver2.sh HUssss_YYYYMMDD_HHMMSS_XXXX.enc HUssss_YYYYMMDD_HHMMSS_XXXX_enckey.enc",
    "The logdecrypt_Ver2.sh tool creates the folder HUssss_YYYYMMDD_HHMMSS_XXXX on the host"),
   pend("X-m two log snapshots taken in separate iterations + comparison")], None),
]


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    tm = {r["swe1_id"]: r for r in csv.DictReader(
        (DATA / "trace_matrix.tsv").open(encoding="utf-8"), delimiter="\t")}
    l2 = {r["swe1_id"]: r for r in csv.DictReader(
        (DATA / "layer2_assign.tsv").open(encoding="utf-8"), delimiter="\t")}
    bo = {r["swe1_id"]: r for r in csv.DictReader(
        (DATA / "batch_order.tsv").open(encoding="utf-8"), delimiter="\t")}
    # batch 2 之範圍 = CCVR 第二批且非 D。
    # CONVERT_DOC rows are produced by gen_deferred.py (batch 3)
    # —— R-SEC23（SEC-11）改該欄後不排除即誤報（SEC-13 §5-3 自報；SEC-14 §1 定案）。
    scope = {k for k, r in bo.items() if r["ccvr_batch"] == "2"
             and r["disposition"] not in ("D", "CONVERT_DOC")}
    planned = {s[0] for s in SPEC}
    assert planned == scope, f"範圍不符：只在計畫 {planned - scope}；只在 batch_order {scope - planned}"

    # sibling_plan_batch2.tsv
    with (DATA / "sibling_plan_batch2.tsv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["swe1_id", "sibling_no", "lower_half(English)", "axis",
                    "trigger_channel", "pending_expected", "priority", "ee"])
        for swe1, no, lower, axis, pr, _d, steps, vm in SPEC:
            chans = []
            for d, cmd, _e in steps:
                if re.match(r"^\s*PENDING:", d):
                    chans.append("PENDING")
                elif cmd and cmd.startswith("Send CAN:"):
                    chans.append("CAN")
                elif cmd and cmd.startswith("$ adb"):
                    chans.append("ADB")
                elif cmd and cmd.startswith("$ openssl") or (cmd or "").startswith("$ ./"):
                    chans.append("HOST")
                elif cmd is None:
                    chans.append("PHYS")
            pend_y = "Y" if any(re.match(r"^\s*PENDING:", d) for d, _c, _e in steps) or \
                     any(str(e).startswith("PENDING:") for _d, _c, e in steps) else "N"
            w.writerow([swe1, no, lower, axis, "/".join(dict.fromkeys(chans)), pend_y, pr, vm or "-"])

    # 037 Description 首句
    import openpyxl as ox
    d_first = {}
    for doc_id, _comp in g1.btm.SWE1_BOOKS:
        wb2 = ox.load_workbook(g1.btm.only(doc_id, "*.xlsx"), read_only=True, data_only=True)
        for r in list(wb2["Analysis Report"].iter_rows(values_only=True))[8:]:
            a = "" if r[0] is None else str(r[0]).strip()
            b = "" if r[1] is None else str(r[1]).strip()
            if not a and not b:
                continue
            key = a or g1.btm.split_source_ids(b)[0].replace(" ", "")
            d_first[key] = g1.first_sentence(str(r[3] or ""))
        wb2.close()

    spec_ref = g1.spec_ref_values()
    wb = openpyxl.load_workbook(g1.TEMPLATE)
    ws = wb[g1.SHEET]
    counter = Counter({"SAM": 6})          # batch 1 之 SAM 末號為 006，自 007 接續
    rows, pend_rows = [], []
    for n, (swe1, no, lower, axis, pr, design, steps, vm) in enumerate(SPEC):
        group = g1.GROUP[tm[swe1]["component"]]
        counter[g1.ABBR[group]] += 1
        tc_id = f"NR1L-{g1.ABBR[group]}-{counter[g1.ABBR[group]]:03d}"
        proc, er = [], []
        for i, (d, cmd, e) in enumerate(steps, 1):
            proc.append(f"{i}. {d}")
            if cmd:
                proc.append(cmd)
            er.append(f"{i}. {e}")
        ph_rec: list[dict] = []
        if g1.SEC08:                                             # R-SEC21
            proc, er, ph_rec = g1.placeholderise(proc, er)
        remarks = [f"source: 037 {swe1} Requirement Description and Verification Criteria",
                   f"sibling axis: {axis}", f'channel_feasible: {tm[swe1]["channel_feasible"]}']
        for rc in ph_rec:                                        # R-SEC21(d)
            remarks.append(f'asset: {rc["x_token"]} — {rc["original_pending"]}')
        if vm:
            if "Send CAN:" in "\n".join(proc):
                msg, sig, raw, label, dbc, cid = CAN[vm]
                remarks.append(f"CAN signal per VHAL User Guide R5 section 2.4; "
                               f"{msg} is {cid}; raw and label verbatim from {dbc}")
        # R-SEC16 amend：六列已入條文，類推註移除。
        r = {"req_id": swe1, "tc_id": tc_id, "test_group": group, "test_set": l2[swe1]["test_set"],
             "test_item": f"{d_first[swe1]}\n({lower})",
             "pre": "\n".join(f"{i}. {x}" for i, x in enumerate(
                 [g1.ADB_ROOT] + (["SAM dongle is inserted into the DUT"] if group == "SAM" else [])
                 + (["Logdog is enabled on the DUT", "Private keys are installed on the DUT",
                     "A USB drive containing __diagSpecialFile and __getSnapshot is inserted into the HU USB port"]
                    if group == "Log Encrypt" else []), 1)),
             "input": "NA", "proc": "\n".join(proc), "er": "\n".join(er),
             "spec": spec_ref[swe1],
             "tc_ref": "NEW", "priority": pr, "design": design, "fs": "No",
             "author": "PeiPYHsu", "remarks": "\n".join(remarks)}
        vmv = VM_HI if vm == "hi" else VM_MI if vm == "mi" else g1.VM_VALUES
        row = g1.FIRST_ROW + n
        for key, col in g1.COLS.items():
            ws[f"{col}{row}"] = r[key]
        for col, val in zip(g1.VM_COLS, vmv):
            ws[f"{col}{row}"] = val
        (OUT_DIR / f"{tc_id}.json").write_text(
            json.dumps({**r, "vehicle_model": dict(zip(g1.VM_NAMES, vmv)),
                        "placeholders": ph_rec},
                       ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        rows.append(r)
        for line in (r["proc"] + "\n" + r["er"]).splitlines():
            m = re.match(r"^\s*(?:\d+\.\s*)?PENDING:\s*(\S+)", line)
            if m:
                pend_rows.append((m.group(1), tc_id))
    report = surgical_save(wb, g1.TEMPLATE, OUT_XLSX)
    print(f"batch2：{len(rows)} TC → {OUT_XLSX.relative_to(ROOT)}")
    print("  逐組:", dict(Counter(r["test_group"] for r in rows)))
    print("  priority:", dict(Counter(r["priority"] for r in rows)))
    print(f"  PENDING 行 {len(pend_rows)}")
    print("  surgical:", {k: v for k, v in report.items() if k != "members_patched"})
    return 0


if __name__ == "__main__":
    sys.exit(main())
