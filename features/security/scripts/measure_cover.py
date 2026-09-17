#!/usr/bin/env python3
"""SEC-06 §3.1 —— 封面類 sheet 與 Reference sheet 之填值量測（唯讀）。

母體：SWC 0708（R-1 v2 基準本）與 Home 0809。兩本不一致者以 SWC 0708 為準。
產物：`features/security/data/cover_fields.tsv`
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "features" / "security" / "data" / "cover_fields.tsv"
BOOKS = {
    "swc": "/Users/peihe/Work/02_Project_R1LR/10_Reviewing/00_TestCase/ASW-R1/SWC/"
           "FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & "
           "Result_SWQT_SWC_20260708.xlsx",
    "home": str(ROOT / "archive/forms_superseded/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT "
                       "STLA Test Case Specification & Result_SWQT_Home_20260809.xlsx"),
}
SHEETS = ["Cover 封面", "ChangeHistory 修訂履歷", "Product Document 記錄封面頁", "Reference"]


def filled(path: str) -> dict[tuple[str, str], str]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for sn in SHEETS:
        if sn not in wb.sheetnames:
            out[(sn, "—")] = "(sheet 不存在)"
            continue
        ws = wb[sn]
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    out[(sn, c.coordinate)] = str(c.value).strip().replace("\n", " ⏎ ")
    wb.close()
    return out


def main() -> int:
    data = {k: filled(v) for k, v in BOOKS.items()}
    keys = sorted(set(data["swc"]) | set(data["home"]),
                  key=lambda k: (SHEETS.index(k[0]) if k[0] in SHEETS else 9, k[1]))
    rows = []
    for k in keys:
        rows.append([k[0], k[1], data["swc"].get(k, ""), data["home"].get(k, ""), ""])
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["sheet", "cell", "swc_value", "home_value", "security_value(proposed)"])
        w.writerows(rows)
    from collections import Counter
    c = Counter(r[0] for r in rows)
    print(f"cover_fields.tsv {len(rows)} 列")
    for sn in SHEETS:
        s = sum(1 for r in rows if r[0] == sn and r[2])
        h = sum(1 for r in rows if r[0] == sn and r[3])
        print(f"  {sn:34s} SWC 填值格 {s:3d}   Home 填值格 {h:3d}")
    print(f"  檔名式：")
    for k, v in BOOKS.items():
        print(f"    {k}: {Path(v).name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
