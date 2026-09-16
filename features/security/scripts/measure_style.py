#!/usr/bin/env python3
"""SEC-03 §1 —— 欄位形制量測（SWC 0708 ＋ pm_29），唯讀。

規則（§1）：Security 之寫法 = 兩本交付本之**多數值**；兩本不一致者以 SWC 0708 為準。
"""
from __future__ import annotations

import collections
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
BOOKS = [
    ("SWC 0708", "/Users/peihe/Work/02_Project_R1LR/10_Reviewing/00_TestCase/ASW-R1/SWC/"
                 "FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & "
                 "Result_SWQT_SWC_20260708.xlsx"),
    ("pm_29", str(ROOT / "features/power/delivered/pm_29.xlsx")),
]
SHEETS = ("Test Case Specification 測試用例規範", "Test Case Specification&Result",
          "Test Case Specification & Result")
HDR = 9
# 0-based；欄序同 R-G1 母本（feature.yaml workbook.columns）
COL = {"req_id": 3, "tc_id": 5, "test_item": 8, "pre": 9, "input": 10,
       "proc": 11, "er": 12, "spec": 13, "tc_ref": 14, "priority": 15,
       "est_time": 16, "design": 17}


def sheet_of(wb):
    for n in SHEETS:
        if n in wb.sheetnames:
            return wb[n]
    for n in wb.sheetnames:
        if "Test Case" in n:
            return wb[n]
    return None


def txt(v) -> str:
    return "" if v is None else str(v).strip()


def main() -> int:
    out = {}
    for label, path in BOOKS:
        p = Path(path)
        if not p.exists():
            print(f"[{label}] 檔不存在：{path}")
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = sheet_of(wb)
        grid = list(ws.iter_rows(values_only=True))
        rows = [r for r in grid[HDR:]
                if len(r) > COL["er"] and any(txt(r[COL[k]]) for k in ("test_item", "proc", "er"))]
        c: dict[str, collections.Counter] = {k: collections.Counter() for k in
                                             ("priority", "design", "input", "tc_ref", "est_time")}
        quotes = collections.Counter()
        for r in rows:
            c["priority"][txt(r[COL["priority"]]) or "(空)"] += 1
            c["design"][txt(r[COL["design"]]) or "(空)"] += 1
            v = txt(r[COL["input"]])
            c["input"]["NA" if v.upper() == "NA" else ("(空)" if not v else "(其他)")] += 1
            c["tc_ref"][txt(r[COL["tc_ref"]]) or "(空)"] += 1
            e = txt(r[COL["est_time"]])
            c["est_time"]["(空)" if not e else ("數值" if re.fullmatch(r"\d+(\.\d+)?", e) else "(其他)")] += 1
            er = txt(r[COL["er"]])
            quotes['"…"'] += len(re.findall(r'"[^"]{1,80}"', er))
            quotes["'…'"] += len(re.findall(r"'[^']{1,80}'", er))
            quotes["`…`"] += len(re.findall(r"`[^`]{1,80}`", er))
        out[label] = (len(rows), c, quotes,
                      [txt(r[COL["test_item"]]) for r in rows[:20]],
                      [txt(r[COL["req_id"]]) for r in rows[:20]])
        wb.close()

    for label, (n, c, q, items, reqs) in out.items():
        print(f"\n{'='*66}\n{label} —— 資料列 {n}")
        for key in ("priority", "design", "input", "tc_ref", "est_time"):
            top = c[key].most_common(6)
            print(f"  {key:10s} {[(k[:46], v) for k, v in top]}")
        print(f"  ER 引號     {dict(q)}")
        print("  test_item 上半（前 5 列抽樣）：")
        for it in items[:5]:
            print("     ", it.splitlines()[0][:96] if it else "(空)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
