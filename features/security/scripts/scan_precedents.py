#!/usr/bin/env python3
"""兩件查證（下放包 _B §3 ＋ _E §3），唯讀。

1. **DID／RID 步驟前例**（R-G13：查無亦回報）—— 掃既有交付本之 Procedure／ER／Pre-Condition，
   找 UDS 讀寫之既有寫法。有則逐字回報出處（檔、sheet、row、原句）。
2. **lint `C`（channel）之假陽性率** —— 以 _E §1(a)(c) 之判準對既有交付本試跑。
   既有交付本**不受 R-SEC{live+6} 拘束**，故其 FAIL 一律為假陽性；
   本數字之用途是「若把 `C` 套到既有語料，會吵成什麼樣」。

用法：python features/security/scripts/scan_precedents.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

SHEET_CANDIDATES = ("Test Case Specification 測試用例規範",
                    "Test Case Specification&Result",
                    "Test Case Specification & Result")
HEADER_ROW = 9
COL = {"test_item": 8, "pre": 9, "input": 10, "proc": 11, "er": 12}  # 0-based: I J K L M

TARGETS = [
    ("SWC 0708（R-1 v2 基準本）", "/Users/peihe/Work/02_Project_R1LR/10_Reviewing/00_TestCase/ASW-R1/SWC/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_SWC_20260708.xlsx"),
    ("DealerMode 20260417(done)", "/Users/peihe/Work/02_Project_R1LR/10_Reviewing/00_TestCase/DealerMode/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_CFTS012_DealerMode_20260417(done).xlsx"),
    ("power pm_29", "features/power/delivered/pm_29.xlsx"),
    ("power pm_73", "features/power/delivered/pm_73.xlsx"),
    ("sw_update 20260830", "features/sw_update/delivered/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_SWUpdate_20260830.xlsx"),
    ("ics_management 20260830", "features/ics_management/delivered/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_ICSManagement_20260830.xlsx"),
    ("vehicle_setting CFTS044", "features/vehicle_setting/delivered/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_CFTS044_Vehicle Controls_20260908.xlsx"),
    ("vsm_v42 20260902", "features/vsm_v42/delivered/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_VehicleSetupManagementR1Low_20260902.xlsx"),
    ("popup 20260908", "features/popup/delivered/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case Specification & Result_SWQT_Popup_20260908.xlsx"),
]

# 1. DID／RID 前例之判準
RE_UDS = re.compile(
    r"\b(?:22|2E|31|10|27|19|3E)\s+[0-9A-Fa-f]{2}\s+[0-9A-Fa-f]{2}\b"
    r"|ReadDataByIdentifier|WriteDataByIdentifier|RoutineControl|DiagnosticSessionControl"
    r"|SecurityAccess|\bDID\b|\bRID\b|\bNRC\b|\bUDS\b|tester\s+present")

# 2. lint C 之判準（_E §1(a)(c)）
RE_STEP_NO = re.compile(r"^\s*(\d+)[.)]\s*(.+)$")
RE_CMD = re.compile(r"^\s*\$\s+\S")
RE_PHYS = re.compile(r"^\s*(?:Insert|Press|Power cycle|Disconnect|Select\s+\")")
RE_OBSERVE = re.compile(
    r"\badb logcat\b|\bod -t x1\b|\badb shell (?:cat|ls)\b"
    r"|Positive response is received|Negative response is received"
    r"|OK \(\d+ test|FAILURES!!!"
    r"|\bopenssl\b|: OK\b|error \d+ at \d+ depth"
    r"|is sent\b|screen is displayed\b")


def sheet_of(wb):
    for n in SHEET_CANDIDATES:
        if n in wb.sheetnames:
            return wb[n]
    for n in wb.sheetnames:
        if "Test Case" in n:
            return wb[n]
    return None


def main() -> int:
    print("=" * 74)
    print("一、DID／RID 步驟前例查證（R-G13：查無亦記明已查）")
    print("=" * 74)
    hits_total = 0
    lintc = []
    for label, path in TARGETS:
        p = Path(path)
        if not p.exists():
            print(f"[{label}] **檔不存在**：{path}")
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = sheet_of(wb)
        if ws is None:
            print(f"[{label}] 無 Test Case sheet（sheetnames={wb.sheetnames[:4]}）")
            wb.close()
            continue
        grid = list(ws.iter_rows(values_only=True))
        rows = [(n, r) for n, r in enumerate(grid[HEADER_ROW:], start=HEADER_ROW + 1)
                if len(r) > COL["er"] and any(r[i] for i in COL.values())]
        hits = []
        fail_proc = fail_er = tested = 0
        for n, r in rows:
            blob = " ".join(str(r[i] or "") for i in COL.values())
            if RE_UDS.search(blob):
                hits.append((n, blob))
            # lint C 試跑
            proc = str(r[COL["proc"]] or "")
            er = str(r[COL["er"]] or "")
            if not proc.strip():
                continue
            tested += 1
            lines = [x for x in proc.splitlines() if x.strip()]
            bad = False
            for i, line in enumerate(lines):
                m = RE_STEP_NO.match(line)
                if not m:
                    continue
                body = m.group(2)
                nxt = lines[i + 1] if i + 1 < len(lines) else ""
                if not (RE_CMD.match(nxt) or RE_PHYS.match(body)):
                    bad = True
                    break
            if bad:
                fail_proc += 1
            if er.strip() and not any(RE_OBSERVE.search(x) for x in er.splitlines() if x.strip()):
                fail_er += 1
        hits_total += len(hits)
        print(f"\n[{label}] 資料列 {len(rows)}；UDS/DID 樣式命中 **{len(hits)}** 列")
        for n, blob in hits[:4]:
            flat = re.sub(r"\s+", " ", blob)[:180]
            print(f"   row {n}: {flat}")
        lintc.append((label, tested, fail_proc, fail_er))
        wb.close()

    print(f"\n→ 九本合計 UDS/DID 樣式命中 **{hits_total}** 列")
    print()
    print("=" * 74)
    print("二、lint `C`（channel）對既有交付本之假陽性率（_E §3）")
    print("=" * 74)
    print("| 交付本 | 有 Procedure 之列 | Procedure FAIL | ER FAIL | Procedure 假陽性率 | ER 假陽性率 |")
    print("|---|---:|---:|---:|---:|---:|")
    T = P = E = 0
    for label, t, fp, fe in lintc:
        T += t; P += fp; E += fe
        print(f"| {label} | {t} | {fp} | {fe} | {fp/t*100:.1f}% | {fe/t*100:.1f}% |" if t else
              f"| {label} | 0 | — | — | — | — |")
    if T:
        print(f"| **合計** | **{T}** | **{P}** | **{E}** | **{P/T*100:.1f}%** | **{E/T*100:.1f}%** |")
    return 0


if __name__ == "__main__":
    sys.exit(main())
