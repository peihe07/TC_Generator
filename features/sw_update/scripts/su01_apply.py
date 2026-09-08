#!/usr/bin/env python3
"""SU-01 §三／§四 —— SW Update 可獨立改善之部分（不待上游）。

來源：`docs/fw036/handoff/down/20260908_SU-01.md`（Pei 指示 2026-09-08）。

**只做兩件**：
  §三  Final Step > 18 詞之縮短（現算全簿取清單，不採下放包之抽樣列舉）
  §四  row 200 之 `'Accept'` → `"Accept"`（proc ＋ ER 兩處）

**不碰**：任何 PENDING、任何 DR、D 欄追溯值、A 類之查證（→ SU-02）。
母本不改（R-G72），落檔一律 `surgical_save`，全域無 `wb.save()`（R16／R-G3）。

§三 之硬規：只動 `test_procedure` 之最末步驟，**ER 一律不動**；
被移除之細節若 ER 未承載，該列**跳過並回報**（不得為縮短而丟失驗證條件）。
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))
from xlsx_surgical import surgical_save  # noqa: E402

warnings.filterwarnings("ignore")

FEAT = ROOT / "features/sw_update"
MASTER = Path("/Users/peihe/Work/02_Project_R1LR/10_Reviewing/00_TestCase/ASW-R2/"
              "SW Update via USB/FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT "
              "STLA Test Case Specification & Result_SWQT_SWUpdate_20260830.xlsx")
MASTER_SHA16 = "d0000d1d96e66f0c"
SB = FEAT / "sandbox/su01"
WORK = SB / "swupdate_20260830.xlsx"
OUT = SB / "swupdate_20260830_Revise1.xlsx"
REPORT = FEAT / "reports/su01_finalstep.tsv"
SHEET = "Test Case Specification 測試用例規範"

C = {"D": 4, "F": 6, "I": 9, "J": 10, "K": 11, "L": 12, "M": 13,
     "N": 14, "P": 16, "R": 18, "AH": 34}
AUTHOR = ("J", "K", "L", "M")
WORD_LIMIT = 18

# 出貨 gate 之基線（下放包 §六.4，執行層落檔前後皆須相符）
DR_BASELINE = {"DR-SU2": 554, "DR-SU6": 52, "DR-SU3": 37, "DR-SU7": 34,
               "DR-SU4": 27, "DR-SU1": 5, "DR-SU5": 3}
PENDING_BASELINE = 195

# ------------------------------------------------------------------ §三 之表
# 一律**純前綴截斷**（不改寫語序）。落檔前逐列 assert：
#   (a) 縮短後 ≤ 18 詞  (b) 仍含 `check that`  (c) 為原句之前綴
#   (d) 被移除之詞彙為同列 ER 末行所承載
NEW_FINAL = {
    10: "Check that the download progress shown on the head unit does not advance",
    11: "Check that the head unit displays a pop-up notification prompting the user "
        "to connect to a Wi-Fi network",
    13: "Check that the head unit Wi-Fi hotspot is available to the companion device again",
    14: "Check that the head unit displays the Wi-Fi pop-up notification after the "
        "seventh failed attempt",
    25: "Check that the head unit offers no control to reject the update",
    30: "Check that the recorded screen content shows the critical update being installed",
    35: "Check that the displayed prompt names Wi-Fi Hotspot, Android Auto and Apple CarPlay",
    53: "Check that the head unit connects to the access point that is within range",
    58: "Check that the head unit displays the pop-up at the ignition off transition",
    60: "Check that no software download over Wi-Fi starts",
    61: "Check that the head unit connects to the saved access point",
    65: "Check that the recorded screen content shows the update being deployed from "
        "the OTA Server",
    69: "Check that the head unit does not start the installation",
    78: "Check that the head unit does not start the installation after the pop-up has closed",
    85: "Check that the head unit displays the conditions not met pop-up",
    86: "Check that Time_remaining equals the difference between Time_scheduled and Time_now",
    95: "Check that the head unit shows no telematics box module update starting",
    96: "Check that the head unit shows no telematics box module update starting",
    97: "Check that the head unit displays the telematics box module update screen",
    103: "Check that the head unit shows no TBM FOTA pop-up and no TBM FOTA status bar entry",
    110: "Check that the recorded screen content of the first run contains the opt-in screen",
    111: "Check that the SW Update screen shows guidance on how to accept the terms "
         "and conditions",
    112: "Check that both the opt-in screen and the download screen show the release notes text",
    114: 'Check that the head unit shows the deployment package details together with '
         'an "Install" option',
    138: "Check that the head unit displays a warning pop-up",
    145: "Check that Version_after equals Version_initial",
    148: "Check that Version_after differs from Version_initial",
    157: "Check that Version_after differs from Version_initial",
    160: "Check that the head unit displays a message stating that the vehicle software "
         "is up to date",
    185: "Check that the head unit shows no update pop-up and no update progress screen",
    270: "Check that the head unit shows the update session ending without the "
         "installation starting",
    276: "Check that the download progress shown on the head unit advances beyond the "
         "value it held",
    291: "Check that the recorded screen content shows the download progress stopping",
    292: "Check that Version_after equals Version_initial",
    307: "Check that the radio changes station and that the settings menu opens",
    325: "Check that the head unit starts and either shows the installation continuing",
}

# 被移除之語意由 ER 以**不同措辭**承載之列：逐列具名 ER 之對應片語，
# 並列出因措辭差異而在逐詞比對中落空之 token。非放寬檢查 —— 該片語須實際存在於 ER。
ER_PARAPHRASE = {
    53: ("shows no connection to the access point that is switched off",
         {"does", "not", "connect"}),
    110: ("the second run shows no opt-in screen", {"and", "is", "shown"}),
    148: ("contains no SW Update prompt and no progress notification",
          {"appears", "in", "or"}),
    157: ("no progress notification and no confirmation screen", {"or"}),
    61: ("the software download over Wi-Fi starting", {"starts"}),
}

SKIP_FINAL = {
    297: "ER 末行為 `PENDING: DR-SU2` 佔位；proc 末步是該列唯一之非 PENDING 斷言，"
         "縮短即丟失比較對象（head-unit-started session）",
    22: "被移除之 `and the engine auto-stop state is not active` 不在 ER —— "
        "ER 末行只寫 `while the ignition is in accessory`，該條件無他處承載",
    175: "被移除之 `while the package downloads` 不在 ER —— ER 末行只寫 "
         "`shows the download progress as a percentage or as a status indication`",
    43: "全句即『動作 ＋ check that ＋ 主要可觀察結果』，無可移除之細節；"
        "19 詞降至 18 以下只能改寫語序，屬改寫非截斷",
    56: "全句即單一斷言（首個連上之 AP ＝ 訊號格最多者）；"
        "任何截斷都會丟失驗證標的 `the one shown with the most signal bars`",
}


def sha256(p: Path) -> str:
    return subprocess.run(["shasum", "-a", "256", str(p)],
                          capture_output=True, text=True,
                          check=True).stdout.split()[0]


def steps(v) -> list[str]:
    return [x for x in str(v or "").split("\n") if x.strip()]


def bare(s: str) -> str:
    return re.sub(r"^\s*\d+\.\s*", "", s).strip()


def words(s: str) -> int:
    return len(bare(s).split())


def tokens(s: str) -> set[str]:
    return set(re.findall(r"[A-Za-z0-9_%$]+", s.lower()))


def gate(ws) -> dict:
    """出貨 gate 與 §五 各項之現算。"""
    data = [r for r in range(10, ws.max_row + 1)
            if str(ws.cell(r, 6).value or "").strip()]
    allcols = lambda r: "\n".join(str(ws.cell(r, c).value or "")
                                  for c in range(1, 35))
    dr: dict[str, int] = {}
    for r in data:
        for m in re.findall(r"DR-SU\d+", allcols(r)):
            dr[m] = dr.get(m, 0) + 1
    return {
        "資料列": len(data),
        "PENDING 列": sum(1 for r in data if "PENDING" in allcols(r)),
        "DR": dr,
        "尾句號": [r for r in data if any(
            x.strip().endswith((".", "。")) for c in AUTHOR for x in steps(ws.cell(r, C[c]).value))],
        "行首尾空白": [r for r in data if any(
            x != x.strip() for c in AUTHOR + ("I", "N")
            for x in str(ws.cell(r, C[c]).value or "").split("\n"))],
        "全形標點": [r for r in data if any(
            ch in str(ws.cell(r, C[c]).value or "") for c in AUTHOR for ch in "；，。、（）")],
        "Proc↔ER 不等": [r for r in data if len(steps(ws.cell(r, 12).value))
                        != len(steps(ws.cell(r, 13).value))],
        "括號下半缺": [r for r in data if not re.search(
            r"\([^)]{5,}\)\s*$", str(ws.cell(r, 9).value or "").strip(), re.S)],
        "Priority 越域": [r for r in data if str(ws.cell(r, 16).value)
                       not in ("P0", "P1", "P2", "P3")],
        "spec_ref 空": [r for r in data if not str(ws.cell(r, 14).value or "").strip()],
        "單引號": [r for r in data if any(
            re.search(r"'[^']{1,40}'", str(ws.cell(r, C[c]).value or "")) for c in AUTHOR)],
        "FinalStep>18": [r for r in data if steps(ws.cell(r, 12).value)
                         and words(steps(ws.cell(r, 12).value)[-1]) > WORD_LIMIT],
    }


def main() -> int:
    assert MASTER.exists(), f"母本不存在：{MASTER}"
    got = sha256(MASTER)
    assert got[:16] == MASTER_SHA16, f"母本 sha16 {got[:16]} ≠ {MASTER_SHA16}，停手"
    SB.mkdir(parents=True, exist_ok=True)
    shutil.copy2(MASTER, WORK)
    assert sha256(WORK) == got, "沙盒副本 sha 不符，停手"

    wb = openpyxl.load_workbook(WORK)
    ws = wb[SHEET]
    before = gate(ws)

    # ---- 現算全簿之 Final Step > 18 詞（不採下放包之列舉）
    data = [r for r in range(10, ws.max_row + 1)
            if str(ws.cell(r, 6).value or "").strip()]
    over = [r for r in data if steps(ws.cell(r, 12).value)
            and words(steps(ws.cell(r, 12).value)[-1]) > WORD_LIMIT]
    last_pending = [r for r in over if "PENDING" in steps(ws.cell(r, 12).value)[-1]]
    target = [r for r in over if r not in last_pending]
    print(f"Final Step > {WORD_LIMIT} 詞：{len(over)} 列")
    print(f"  末步本身即 PENDING 行（不可縮）：{len(last_pending)} 列 {last_pending}")
    print(f"  可施作母體：{len(target)} 列")
    assert set(target) == set(NEW_FINAL) | set(SKIP_FINAL), (
        f"母體與本檔之表不符：缺 {set(target)-set(NEW_FINAL)-set(SKIP_FINAL)}，"
        f"多 {(set(NEW_FINAL)|set(SKIP_FINAL))-set(target)}")

    # ---- §三 施作
    rows_out = []
    for r in sorted(NEW_FINAL):
        L = steps(ws.cell(r, 12).value)
        E = steps(ws.cell(r, 13).value)
        old, new = bare(L[-1]), NEW_FINAL[r]
        assert old.startswith(new), f"r{r} 非前綴截斷"
        assert len(new.split()) <= WORD_LIMIT, f"r{r} 縮短後仍 {len(new.split())} 詞"
        assert "check that" in new.lower(), f"r{r} 缺 check that"
        missing = tokens(old) - tokens(new) - tokens(E[-1])
        if r in ER_PARAPHRASE:
            phrase, allowed = ER_PARAPHRASE[r]
            assert phrase.lower() in E[-1].lower(), \
                f"r{r} 之 ER 未含所具名之片語：{phrase!r}"
            missing -= allowed
        assert not missing, f"r{r} 被移除之詞未被 ER 承載：{sorted(missing)}"
        num = re.match(r"^\s*(\d+)\.", L[-1]).group(1)
        L[-1] = f"{num}. {new}"
        ws.cell(r, 12).value = "\n".join(L)
        rows_out.append({"row": r, "tc_id": ws.cell(r, 6).value,
                         "words_before": len(old.split()),
                         "words_after": len(new.split()),
                         "action": "shortened",
                         "note": ("ER 以不同措辭承載：" + ER_PARAPHRASE[r][0]
                                  if r in ER_PARAPHRASE else "")})
    for r in sorted(SKIP_FINAL):
        L = steps(ws.cell(r, 12).value)
        rows_out.append({"row": r, "tc_id": ws.cell(r, 6).value,
                         "words_before": words(L[-1]), "words_after": words(L[-1]),
                         "action": "skipped", "note": SKIP_FINAL[r]})

    # ---- §四 施作：row 200 之引號形制（字串不改）
    for col in (12, 13):
        v = str(ws.cell(200, col).value or "")
        assert "'Accept'" in v, f"row 200 c{col} 找不到 'Accept'"
        ws.cell(200, col).value = v.replace("'Accept'", '"Accept"')
    # test_item 上半為 CFTS057 逐字（其原文即單引號），不動
    assert "'Accept'" in str(ws.cell(200, 9).value), "test_item 之 verbatim 應保留單引號"

    after = gate(ws)

    # ---- 落檔前 gate
    assert after["PENDING 列"] == before["PENDING 列"] == PENDING_BASELINE, \
        f"PENDING {before['PENDING 列']} → {after['PENDING 列']}，應為 {PENDING_BASELINE}"
    assert after["DR"] == before["DR"] == DR_BASELINE, \
        f"DR 計數變動：{before['DR']} → {after['DR']}"
    for k in ("資料列", "尾句號", "行首尾空白", "全形標點", "Proc↔ER 不等",
              "括號下半缺", "Priority 越域", "spec_ref 空"):
        assert before[k] == after[k], f"{k} 變動：{before[k]} → {after[k]}"
    assert not after["單引號"], f"單引號未歸零：{after['單引號']}"
    assert set(after["FinalStep>18"]) == set(last_pending) | set(SKIP_FINAL), \
        f"殘留之 >18 列不符預期：{after['FinalStep>18']}"

    report = surgical_save(wb, WORK, OUT)

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    import csv
    with open(REPORT, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, delimiter="\t", fieldnames=[
            "row", "tc_id", "words_before", "words_after", "action", "note"])
        w.writeheader()
        w.writerows(sorted(rows_out, key=lambda x: x["row"]))

    print(f"\n§三 施作 {len(NEW_FINAL)} 列；跳過 {len(SKIP_FINAL)} 列 {sorted(SKIP_FINAL)}")
    print(f"§四 row 200 引號 2 處；test_item 之 verbatim 未動")
    print(f"\nPENDING {before['PENDING 列']} → {after['PENDING 列']}")
    print(f"DR      {after['DR']}")
    print(f"Final Step > 18 詞  {len(before['FinalStep>18'])} → "
          f"{len(after['FinalStep>18'])}（餘為 {len(last_pending)} 列末步即 PENDING "
          f"＋ {len(SKIP_FINAL)} 列具名跳過）")
    print(f"\n母本   sha256 {sha256(MASTER)}")
    print(f"沙盒本 {WORK.relative_to(ROOT)}  sha256 {sha256(WORK)}")
    print(f"Revise1 {OUT.relative_to(ROOT)}  sha256 {sha256(OUT)}")
    print(f"surgical_save: {report}")
    print(f"逐列報表 → {REPORT.relative_to(ROOT)}")
    assert sha256(MASTER)[:16] == MASTER_SHA16, "母本 sha 改變，R-G72 破，停手"
    return 0


if __name__ == "__main__":
    sys.exit(main())
