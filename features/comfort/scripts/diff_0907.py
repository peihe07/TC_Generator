#!/usr/bin/env python3
"""CMF-01 1 節：0907 本 ↔ generated/*.json（經 write_back 之匯出規則）逐格對帳。

唯讀。三方比對：
  A = 0907 本（Pei 自 Reviewing 夾取之工作簿，sha 3550d16d…）
  G = generated/*.json 經 `write_back.row_plan()`／`cell_of()` 所得之「應寫值」
  R = ENTRY 035 之產出檔（…_Comfort_20260817_rowsort.xlsx）

G 與 R 應逐格相同（ENTRY 035 §3.3 已驗）；此處重驗一次，使「A 與 G 之差」
可歸屬為「A 相對 R 之離線編修」而非「語料在 0817 之後前進」。

Usage:
    python3 features/comfort/scripts/diff_0907.py <0907.xlsx> <out.tsv>
"""

import csv
import sys
from collections import Counter
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import write_back as wb_mod  # noqa: E402

ROWSORT = wb_mod.OUT
SHEET = wb_mod.SHEET
FIRST = wb_mod.FIRST_ROW
COLS = wb_mod.COLS
# 未由 pipeline 寫入之欄，亦與 R 比對（辨識 T–Z 等離線改動）
EXTRA = ["B", "C", "E", "O", "Q", "T", "U", "V", "W", "X", "Y", "Z",
         "AA", "AB", "AC", "AD", "AE", "AF", "AG"]


def norm(v) -> str:
    return "" if v is None else str(v)


def main() -> int:
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    plan = wb_mod.row_plan(wb_mod.load_tcs())
    seq = wb_mod.tcid_sequence(plan)
    wa = openpyxl.load_workbook(src)[SHEET]
    wr = openpyxl.load_workbook(ROWSORT)[SHEET]

    last_a = max(r for r in range(FIRST, wa.max_row + 1)
                 if any(norm(wa[f"{c}{r}"].value) for c in "DFI"))
    n_rows = max(last_a - FIRST + 1, len(plan))
    rows, col_ag, col_ar, col_rg, extra_ar = [], Counter(), Counter(), Counter(), Counter()
    for i in range(n_rows):
        r = FIRST + i
        p = plan[i] if i < len(plan) else None
        rec = {"row": r,
               "json_tc_id": (p.get("tc_id", "") if p else "<none>"),
               "json_req_id": (p["req_id"] if p else "<none>")}
        diffs = []
        for col, field in COLS.items():
            a = norm(wa[f"{col}{r}"].value)
            g = wb_mod.cell_of(p, field, seq[i]) if p else ""
            rr = norm(wr[f"{col}{r}"].value)
            if a != g:
                col_ag[col] += 1
                diffs.append(col)
            if a != rr:
                col_ar[col] += 1
            if rr != g:
                col_rg[col] += 1
        xd = []
        for col in EXTRA:
            if norm(wa[f"{col}{r}"].value) != norm(wr[f"{col}{r}"].value):
                extra_ar[col] += 1
                xd.append(col)
        rec.update({"F_0907": norm(wa[f"F{r}"].value),
                    "D_0907": norm(wa[f"D{r}"].value),
                    "n_diff_vs_json": len(diffs),
                    "diff_cols_vs_json": ",".join(diffs),
                    "diff_noncorpus_cols_vs_rowsort": ",".join(xd)})
        rows.append(rec)

    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter="\t")
        w.writeheader()
        w.writerows(rows)
        f.write("\n# summary\n")
        f.write(f"# rows_0907={last_a - FIRST + 1}\trows_plan={len(plan)}\t"
                f"tcs_json={sum(1 for x in plan if not x.get(wb_mod.BLANK))}\n")
        f.write("# col\tA(0907)!=G(json)\tA!=R(rowsort)\tR!=G\n")
        for col in COLS:
            f.write(f"# {col}\t{col_ag[col]}\t{col_ar[col]}\t{col_rg[col]}\n")
        f.write("# non-corpus cols A!=R: " +
                ", ".join(f"{c}={n}" for c, n in extra_ar.items()) + "\n")
        f.write(f"# rows with any corpus-col diff vs json: "
                f"{sum(1 for x in rows if x['n_diff_vs_json'])}\n")
    print(out.read_text(encoding="utf-8").split("# summary")[1])
    return 0


if __name__ == "__main__":
    sys.exit(main())
