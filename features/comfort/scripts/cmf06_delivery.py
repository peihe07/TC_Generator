#!/usr/bin/env python3
"""CMF-06 §3 —— Comfort HMI **交付候選本**（0907 本之複本，xlsx_surgical 寫入）。

內容 ＝ 候選本 v4 之全部變更 ＋ 封面類分頁（Cover／ChangeHistory／Product
Document）＋ `J5` 日期。**車型欄與 TC 內容與 v4 逐格相同** —— CMF-06 §2 所指之
兩條（工作簿 F `-403`／`-404`，即 req `116-03`／`116-04`）實查**已為 V=Y=Z=0**，
無格可改；R-C67 之真實涵蓋面見上繳包 §2，本輪不套用。

**封面類分頁之填值依據**（CMF-06 §3-2「先量兩本同名分頁」）：
`SWC 0708` 交付本**不在本 repo**（只有 SWC 20260628／20260529），故以
`SWC 20260628` 與 `Home 20260809` 兩本，加最近之兩本交付（`Popup 20260908`、
`Comfort 20260907` 自身）共四本量測；逐格結果見上繳包 §3。

**`Test Case Framework` 之 SmartArt 舊名只報不改**（Tier 3，CMF-06 §3-3）。
**新列 476–485 之列高與樣式做不到** —— `backend/xlsx_surgical` 之寫入路徑
依設計只 patch 儲存格**值**（`diff_cells`），列高（`<row ht=…>`）與樣式（`<c s=…>`）
不在其內；另開一條 XML 編修路徑會破壞本 repo 之單一寫入路徑紀律
（`tests/test_single_write_path.py` 所守）。量測值附於上繳包 §3-4，供 Excel 內一次設定。

Usage:
    python3 features/comfort/scripts/cmf06_delivery.py
"""

import csv
import datetime as dt
import hashlib
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import to_excel

sys.path.insert(0, str(Path(__file__).resolve().parent))
import write_back as W                                           # noqa: E402
sys.path.insert(0, str(W.ROOT))
from backend.xlsx_surgical import StructureError, sheet_members, surgical_save  # noqa: E402

FEATURE = W.FEATURE
SRC = FEATURE / "delivered" / (
    "FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case "
    "Specification & Result_SWQT_ComfortHMI_20260907.xlsx")
SRC_SHA = "3550d16d8bd319633d872856ae51ba63a7d2a928d22fab69bff6dd1963736aff"
PREV = FEATURE / "sandbox" / "revise_0922b" / (
    "FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case "
    "Specification & Result_SWQT_ComfortHMI_20260922_v4.xlsx")
PREV_SHA = "7b3912fcaf4780ff7490e76c2fba18bc00881f657521615b656266c3a89d6bd4"
OUT_DIR = FEATURE / "sandbox" / "delivery"
OUT = OUT_DIR / (
    "FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case "
    "Specification & Result_SWQT_ComfortHMI_20260922.xlsx")
REPORTS = FEATURE / "docs" / "reports"
RECON_TSV = REPORTS / "cmf06_delivery_cell_recon.tsv"
PREV_TSV = REPORTS / "cmf06_vs_v4_cell_delta.tsv"
CHANGE_LOGS = tuple(REPORTS / f"cmf0{n}_json_change_log.tsv" for n in (2, 3, 4, 5))
VM_ASSIGN = REPORTS / "vm_assign.tsv"

FIRST, LAST_EXISTING = 10, 475
ROW_019_03 = 116
NEW_ROW_FIRST = 476
AUTHOR, TC_REF = "PeiPYHsu", "NEW"
VM_FIXED = {"T": 0, "U": 0, "W": 0, "X": 0}
VM_PER_TC = {"V": "V_Promaster", "Y": "Y_Toro", "Z": "Z_Fastback"}
VM_COLS = tuple(VM_FIXED) + tuple(VM_PER_TC)
FIELD_OF_COL = {c: f for c, f in W.COLS.items() if c != "F"}

# ------------------------------------------------------------ 封面類分頁
COVER, CHG, PROD = ("Cover 封面", "ChangeHistory 修訂履歷",
                    "Product Document 記錄封面頁")
DELIVERY_DATE = dt.datetime(2026, 9, 22)
# `backend/xlsx_surgical` 之寫入路徑把 `datetime` 序列化為**字串**，落到交付本即成
# 文字型日期（Excel 不認）。故日期格改寫 **Excel 序列值**；該格之 number_format
# 為日期式者，openpyxl 讀回即為 `datetime`，與意圖值相符。
DELIVERY_SERIAL = to_excel(DELIVERY_DATE)   # 46287.0
J5_TEXT = "2026-09-22"                      # J5 為字串欄（number_format `@`）
REVISION = "D"                              # 0907 本之現行版號 C ＋1
PROD_VERSION = "V1.1"                       # Product Document 之 V1.0 ＋1
SUMMARY = (
    "車型欄依 PROXI 逐條填值、DR-46 結案（CAN 訊號改用實際名與值）、"
    "Test Set 名稱修正、新增 sibling 10 條\n"
    "Filled the vehicle model columns per PROXI configuration, closed DR-46 by "
    "using the actual CAN signal names and values, corrected the Test Set "
    "names, and added 10 sibling test cases")
# **留空給 Pei 之三格**（CMF-06 §3-2）—— 於此具名，不寫值，亦不清除既有值
LEFT_BLANK = [
    (CHG, "E8", "核准者 Approver（新列 D 之核准者）"),
    (COVER, "G7", "核准者日期（0907 本即空；Home 0809 有值）"),
    (COVER, "G8", "審查者日期（0907 本即空；Home 0809 有值）"),
]


def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def norm(v):
    return "" if v is None else v


def dv_sqrefs(path: Path) -> list:
    with zipfile.ZipFile(path) as z:
        xml = z.read(sheet_members(path)[W.SHEET]).decode("utf-8")
    return (re.findall(r'<dataValidation[^>]*sqref="([^"]+)"', xml)
            + re.findall(r"<xm:sqref>([^<]+)</xm:sqref>", xml))


def load_vm() -> dict:
    with VM_ASSIGN.open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(
            (l for l in fh if not l.startswith("#")), delimiter="\t"))
    return {r["tc_id"]: {c: int(r[k]) for c, k in VM_PER_TC.items()} for r in rows}


def main() -> int:
    if sha256(SRC) != SRC_SHA:
        raise SystemExit("母本 sha256 不符 ENTRY 036 —— 停")
    if sha256(PREV) != PREV_SHA:
        raise SystemExit("候選本 v4 sha256 不符 ENTRY 040 —— 停")
    vm = load_vm()
    tcs = W.load_tcs()
    existing = [t for t in tcs if int(t["tc_id"][-3:]) <= 466]
    new_tcs = [t for t in tcs if int(t["tc_id"][-3:]) > 466]
    plan = W.row_plan(existing)
    if len(plan) != LAST_EXISTING - FIRST + 1 or any(p.get(W.BLANK) for p in plan):
        raise SystemExit(f"row plan is not 466 TC rows: {len(plan)}")

    src_wb = openpyxl.load_workbook(SRC)
    src = src_wb[W.SHEET]
    wb = openpyxl.load_workbook(SRC)
    ws = wb[W.SHEET]

    intended = {}                            # (sheet, row, col) -> value
    c_of_req = {}
    for i, t in enumerate(plan):
        r = FIRST + i
        if src[f"D{r}"].value != t["req_id"]:
            raise SystemExit(f"row {r}: D={src[f'D{r}'].value!r} != {t['req_id']}")
        c_of_req.setdefault(t["req_id"], src[f"C{r}"].value)
        for col, field in FIELD_OF_COL.items():
            v = W.render(field, t[field])
            intended[(W.SHEET, r, col)] = v if v != "" else None
        if r == ROW_019_03:
            intended[(W.SHEET, r, "F")] = t["tc_id"]
        for col, v in VM_FIXED.items():
            intended[(W.SHEET, r, col)] = v
        for col, v in vm[t["tc_id"]].items():
            intended[(W.SHEET, r, col)] = v
    for j, t in enumerate(sorted(new_tcs, key=lambda x: int(x["tc_id"][-3:]))):
        r = NEW_ROW_FIRST + j
        if src[f"D{r}"].value is not None:
            raise SystemExit(f"row {r} is not empty in the source")
        for col, field in FIELD_OF_COL.items():
            v = W.render(field, t[field])
            intended[(W.SHEET, r, col)] = v if v != "" else None
        intended[(W.SHEET, r, "F")] = t["tc_id"]
        intended[(W.SHEET, r, "C")] = c_of_req[t["req_id"]]
        intended[(W.SHEET, r, "O")] = TC_REF
        intended[(W.SHEET, r, "AA")] = AUTHOR
        for col, v in VM_FIXED.items():
            intended[(W.SHEET, r, col)] = v
        for col, v in vm[t["tc_id"]].items():
            intended[(W.SHEET, r, col)] = v

    # ---- CMF-06 §3-3：J5 ----------------------------------------------------
    intended[(W.SHEET, 5, "J")] = J5_TEXT
    # ---- CMF-06 §3-2：封面類分頁 ---------------------------------------------
    intended[(COVER, 9, "G")] = DELIVERY_DATE            # 修訂日期（fmt yyyy/mm/dd）
    intended[(CHG, 8, "A")] = REVISION                   # 版本 C → D
    intended[(CHG, 8, "B")] = SUMMARY
    intended[(CHG, 8, "C")] = AUTHOR                     # 修訂人
    intended[(CHG, 8, "D")] = DELIVERY_DATE              # fmt yyyy/mm/dd
    intended[(PROD, 8, "B")] = DELIVERY_DATE             # 日期（fmt mm-dd-yy）
    intended[(PROD, 14, "A")] = PROD_VERSION
    intended[(PROD, 14, "B")] = SUMMARY
    intended[(PROD, 14, "C")] = "許珮瑜 PeiPYHsu"        # 與 C13 同形
    # `Product Document`!D14 之 number_format 為 `General`（範本空白列），非日期式 ——
    # 寫序列值會顯示成 46287。故此格寫**字串**，與 D13（`m"月"d"日"`）之呈現不同，
    # 於上繳包具名，由 Pei 於 Excel 內設為日期格式即可。
    intended[(PROD, 14, "D")] = J5_TEXT
    # 寫入時之實值：日期格改寫 Excel 序列值（見 DELIVERY_SERIAL 之註）
    write_as = {(COVER, 9, "G"): DELIVERY_SERIAL,
                (CHG, 8, "D"): DELIVERY_SERIAL,
                (PROD, 8, "B"): DELIVERY_SERIAL}
    for sheet, coord, _why in LEFT_BLANK:                # 明示不寫、亦不清除
        col = re.match(r"([A-Z]+)(\d+)", coord).group(1)
        row = int(re.match(r"([A-Z]+)(\d+)", coord).group(2))
        if src_wb[sheet][coord].value is not None:
            raise SystemExit(f"{sheet}!{coord} 於母本非空，留空之前提不成立")
        intended.pop((sheet, row, col), None)

    for (sheet, r, col), v in intended.items():
        wb[sheet][f"{col}{r}"] = write_as.get((sheet, r, col), v)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        report = surgical_save(wb, SRC, OUT)
    except StructureError as exc:
        print(f"ABORTED (structure invariant): {exc}", file=sys.stderr)
        return 1

    ok = True

    def g(name, passed, note=""):
        nonlocal ok
        ok &= passed
        print(f"- {'PASS' if passed else '**FAIL**'} — {name}" + (f" — {note}" if note else ""))

    print("## §3 交付候選本 —— invariant 與逐格對帳（自產出檔讀回）\n")
    with zipfile.ZipFile(SRC) as a, zipfile.ZipFile(OUT) as b:
        g("zip member 數與母本相同（ABORT 級）", len(a.namelist()) == len(b.namelist()),
          f"{len(a.namelist())} vs {len(b.namelist())}")
    mem = sheet_members(SRC)
    expect = sorted({mem[W.SHEET], mem[COVER], mem[CHG], mem[PROD]})
    g("差異僅限本包所改之四個 sheet 之 xml", sorted(report["differing"]) == expect,
      f"{sorted(report['differing'])} vs {expect}")
    s_src, s_out = dv_sqrefs(SRC), dv_sqrefs(OUT)
    g("DV sqref（classic ＋ x14）逐字相同（ABORT 級）", s_src == s_out, f"{s_out}")

    out_wb = openpyxl.load_workbook(OUT)
    out = out_wb[W.SHEET]
    last_row = NEW_ROW_FIRST + len(new_tcs) - 1
    bad_b = [r for r in range(FIRST, 602)
             if out[f"B{r}"].value != f'=IF(ISBLANK($D{r}),"",ROW()-9)']
    g(f"B 欄公式 row {FIRST}–601 逐列原樣（涵蓋新列 {NEW_ROW_FIRST}–{last_row}）", not bad_b,
      str(bad_b[:5]))
    blanks = [(s, c) for s, c, _ in LEFT_BLANK if out_wb[s][c].value is not None]
    g("留空之三格於產出檔仍為空", not blanks, str(blanks))

    # ---- 逐格對帳：對 0907 本，全部工作表 ---------------------------------------
    changed_fields = {}
    for path in CHANGE_LOGS:
        with path.open(encoding="utf-8") as fh:
            for row in csv.DictReader(fh, delimiter="\t"):
                changed_fields.setdefault((row["tc_id"], row["field"]), set()).add(row["item"])
    tc_at_row = {FIRST + i: t for i, t in enumerate(plan)}
    tc_at_row.update({NEW_ROW_FIRST + j: t for j, t in enumerate(
        sorted(new_tcs, key=lambda x: int(x["tc_id"][-3:])))})

    def attribution(sheet, r, col):
        if sheet in (COVER, CHG, PROD):
            return "CMF-06 §3-2 封面類分頁"
        if sheet != W.SHEET:
            return ""
        if (r, col) == (5, "J"):
            return "CMF-06 §3-3 J5"
        if r < FIRST:
            return ""
        if r >= NEW_ROW_FIRST:
            return "§3 sibling"
        if col in VM_COLS:
            return "R-C48(amend2)/R-C62 車型欄"
        if r == ROW_019_03:
            return "R-C51"
        if col == "N":
            return "R-C49(amend)"
        t = tc_at_row[r]
        items = changed_fields.get((t["tc_id"], FIELD_OF_COL.get(col, "")), set())
        return "+".join(sorted(items))

    recon, mismatch, wrong, n_cells = [], [], [], 0
    for name in src_wb.sheetnames:
        a, b = src_wb[name], out_wb[name]
        rows = max(a.max_row, b.max_row)
        cols = max(a.max_column, b.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                n_cells += 1
                va, vb = a.cell(r, c).value, b.cell(r, c).value
                col = openpyxl.utils.get_column_letter(c)
                key = (name, r, col)
                if key in intended:
                    if norm(vb) != norm(intended[key]):
                        wrong.append(f"{name}!{col}{r}")
                    if norm(va) != norm(vb):
                        item = attribution(name, r, col)
                        if not item:
                            mismatch.append(f"{name}!{col}{r} (unattributed)")
                        recon.append({"sheet": name, "cell": f"{col}{r}",
                                      "tc_id": tc_at_row.get(r, {}).get("tc_id", "")
                                      if name == W.SHEET else "",
                                      "item": item or "UNATTRIBUTED",
                                      "before": "" if va is None else str(va),
                                      "after": "" if vb is None else str(vb)})
                elif norm(va) != norm(vb):
                    mismatch.append(f"{name}!{col}{r}")
    g("寫入格讀回 == 意圖值", not wrong, f"{len(wrong)} 格不符 {wrong[:5]}")
    g("**未列於本包之格與 0907 本逐格相同**（全部工作表）", not mismatch,
      f"不符 {len(mismatch)} 格 {mismatch[:5]}；比對 {n_cells} 格")
    by_item = {}
    for x in recon:
        by_item[x["item"]] = by_item.get(x["item"], 0) + 1
    print(f"- MEASURED — 對 0907 本之變更格 {len(recon)}：" +
          "、".join(f"{k} {v}" for k, v in sorted(by_item.items())))
    with RECON_TSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["sheet", "cell", "tc_id", "item",
                                           "before", "after"], delimiter="\t")
        w.writeheader()
        w.writerows(recon)

    # ---- 逐格對帳：對候選本 v4 -------------------------------------------------
    prev_wb = openpyxl.load_workbook(PREV)
    delta, unattr = [], []
    for name in prev_wb.sheetnames:
        a, b = prev_wb[name], out_wb[name]
        rows = max(a.max_row, b.max_row)
        cols = max(a.max_column, b.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                va, vb = a.cell(r, c).value, b.cell(r, c).value
                if norm(va) == norm(vb):
                    continue
                col = openpyxl.utils.get_column_letter(c)
                item = attribution(name, r, col)
                if not item:
                    unattr.append(f"{name}!{col}{r}")
                delta.append({"sheet": name, "cell": f"{col}{r}",
                              "item": item or "UNATTRIBUTED",
                              "v4": "" if va is None else str(va),
                              "delivery": "" if vb is None else str(vb)})
    g("相對候選本 v4 之每一差異格皆歸因到本包一項", not unattr,
      f"未歸因 {len(unattr)} 格 {unattr[:5]}")
    by2 = {}
    for x in delta:
        by2[x["item"]] = by2.get(x["item"], 0) + 1
    print(f"- MEASURED — 相對 v4 之差異格 {len(delta)}：" +
          "、".join(f"{k} {v}" for k, v in sorted(by2.items())))
    with PREV_TSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["sheet", "cell", "item", "v4", "delivery"],
                           delimiter="\t")
        w.writeheader()
        w.writerows(delta)
    print(f"- recon tables: {RECON_TSV.relative_to(FEATURE)}, "
          f"{PREV_TSV.relative_to(FEATURE)}")
    print("- 留空給 Pei：" + "、".join(f"{s}!{c}（{why}）" for s, c, why in LEFT_BLANK))
    print(f"\noutput: {OUT.relative_to(FEATURE)}\nsha256: {sha256(OUT)}")
    print(f"MISMATCH_COUNT: {len(mismatch)}")
    print(f"TC_TOTAL: {len(tcs)}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
