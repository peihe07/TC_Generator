#!/usr/bin/env python3
"""PU-02 §二／§三 —— Popup 交付：封面補齊 ＋ 檔案搬移 ＋ MANIFEST。

來源：`docs/fw036/handoff/down/20260908_PU-02.md`（Pei 裁 2026-09-08）。

輸入 `output/FM-WI-FSM-036-A01 …_Popup_20260828.xlsx`（sha16 a13559d51d91d369，
與 `10_Reviewing/` 同名檔逐位元相同）。
**`Test Case Specification 測試用例規範` 分頁一格不動**；只寫
`Cover 封面` 與 `Product Document 記錄封面頁`。
`ChangeHistory 修訂履歷` 依 §2.2 **不寫**。

落檔一律 `backend.xlsx_surgical.surgical_save()`，全域無 `wb.save()`（R16／R-G3）。
"""

from __future__ import annotations

import csv
import re
import shutil
import subprocess
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))
from xlsx_surgical import surgical_save  # noqa: E402

warnings.filterwarnings("ignore")

FEAT = ROOT / "features/popup"
OUT_DIR = FEAT / "output"
DELIVERED = FEAT / "delivered"
SRC = OUT_DIR / ("FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case "
                 "Specification & Result_SWQT_Popup_20260828.xlsx")
SRC_SHA16 = "a13559d51d91d369"
STAGE = FEAT / "sandbox/_pu02_stage.xlsx"

DELIVER_NAME = ("FM-WI-FSM-036-A01 STLA 測試用例規範與結果_SWQT STLA Test Case "
                "Specification & Result_SWQT_Popup_20260908.xlsx")

TC_SHEET = "Test Case Specification 測試用例規範"
COVER = "Cover 封面"
PRODDOC = "Product Document 記錄封面頁"
CHANGEHIST = "ChangeHistory 修訂履歷"

AUTHOR = "PeiPYHsu"
REV_DATE = "2026-09-08"

# §二 之補欄。位址皆由合併範圍實測而定（見上繳 §4）。
# B3／B5／B6／B7 為跨簿套用（值逐字取自 CFTS044 交付本同欄，Pei 准 2026-09-08）；
# B4 為本簿自有（§2.3 之表指定為本簿交付檔名，非沿用 CFTS044 之截斷字串）。
EDITS = {
    COVER: {"D9": AUTHOR, "G9": REV_DATE},
    PRODDOC: {"B3": "new R1L",
              "B4": DELIVER_NAME,
              "B5": "V1.0",
              "B6": "SW Testing",
              "B7": "Confidential",
              "B8": REV_DATE,
              "A13": "V1.0",
              "B13": "Initial Release 初版發布",
              "C13": AUTHOR,
              "D13": REV_DATE},
}

NOTE = ("Popup v1；5 條 TC／5 Functional leaf（037 V0.2 全覆蓋）；"
        "Pop Up List 基線 SR24 Post 2A (Dec 15, 2023)；"
        "Core §5.1–5.4 缺口見 COVERAGE_GAPS.md")


def sha256(p: Path) -> str:
    return subprocess.run(["shasum", "-a", "256", str(p)],
                          capture_output=True, text=True,
                          check=True).stdout.split()[0]


def sheet_map(p: Path) -> dict[str, str]:
    """分頁名 -> zip member，經 workbook.xml.rels 解析（不臆測編號）。"""
    z = zipfile.ZipFile(p)
    rid = {m.group(1): m.group(2) for m in re.finditer(
        r'Id="([^"]+)"[^>]*Target="([^"]+)"',
        z.read("xl/_rels/workbook.xml.rels").decode())}
    return {m.group(1): "xl/" + rid[m.group(2)].lstrip("/")
            for m in re.finditer(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"',
                                 z.read("xl/workbook.xml").decode())}


def merged_of(ws, coord: str) -> str:
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return str(rng)
    return "(未合併)"


def lines(v) -> list[str]:
    return [x for x in str(v or "").split("\n") if x.strip()]


def main() -> int:
    assert SRC.exists(), f"來源不存在：{SRC}"
    got = sha256(SRC)
    assert got[:16] == SRC_SHA16, f"來源 sha16 {got[:16]} ≠ {SRC_SHA16}，停手"

    probe = subprocess.run(
        ["git", "check-ignore", "-v", str(DELIVERED / DELIVER_NAME)],
        cwd=ROOT, capture_output=True, text=True)
    assert probe.returncode != 0, (
        f"delivered/ 被 gitignore 命中，停手：{probe.stdout.strip()}")

    smap = sheet_map(SRC)
    print("分頁 → zip member（workbook.xml.rels 實測）")
    for k in (COVER, CHANGEHIST, PRODDOC, TC_SHEET):
        print(f"  {k:34} {smap[k]}")

    STAGE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, STAGE)
    wb = openpyxl.load_workbook(STAGE)
    print("\n合併範圍實測（§2.1 之 ⚠）")
    for sheet, cells in EDITS.items():
        ws = wb[sheet]
        for coord, val in cells.items():
            before = ws[coord].value
            print(f"  {sheet[:12]:12} {coord:4} merged={merged_of(ws, coord):10} "
                  f"before={before!r}")
            assert before in (None, ""), f"{sheet}!{coord} 非空（{before!r}），停手"
            ws[coord] = val

    out_path = OUT_DIR / DELIVER_NAME
    report = surgical_save(wb, STAGE, out_path)
    STAGE.unlink()

    # ------------------------------------------------ §四 交付前最後驗證
    print("\n== §四 交付前最後驗證 ==")
    za, zb = zipfile.ZipFile(SRC), zipfile.ZipFile(out_path)
    na, nb = set(za.namelist()), set(zb.namelist())
    assert na == nb, f"zip member 集合相異：{na ^ nb}"
    differing = sorted(n for n in na if za.read(n) != zb.read(n))
    want = sorted({smap[COVER], smap[PRODDOC]})
    print(f"  1. member 數 {len(na)} → {len(nb)}；相異 {differing}")
    assert differing == want, f"相異 member 應僅 {want}，實得 {differing}"
    for tag, member in (("TC", smap[TC_SHEET]), ("ChangeHistory", smap[CHANGEHIST])):
        assert za.read(member) == zb.read(member), f"{tag} xml 有 diff，停手"
        print(f"     {tag} xml diff = 0")
    dv = (za.read(smap[TC_SHEET]).decode().count("<dataValidation "),
          zb.read(smap[TC_SHEET]).decode().count("<dataValidation "))
    assert dv[0] == dv[1], f"dv 計數 {dv}"
    print(f"     dataValidation {dv[0]} → {dv[1]}")

    wa = openpyxl.load_workbook(SRC)[TC_SHEET]
    ws = openpyxl.load_workbook(out_path)[TC_SHEET]
    diff = [(r, c) for r in range(1, wa.max_row + 1) for c in range(1, 35)
            if wa.cell(r, c).value != ws.cell(r, c).value]
    assert not diff, f"TC 分頁有 {len(diff)} 格相異，停手：{diff[:5]}"
    print(f"  2. TC 分頁逐格 diff = 0（{wa.max_row} × 34）")

    C = {"D": 4, "F": 6, "I": 9, "J": 10, "K": 11, "L": 12, "M": 13,
         "N": 14, "P": 16, "R": 18, "AH": 34}
    A = ("J", "K", "L", "M")
    data = [r for r in range(10, ws.max_row + 1)
            if str(ws.cell(r, 6).value or "").strip()]
    blob = lambda r: "\n".join(str(ws.cell(r, C[k]).value or "")
                               for k in ("I", "J", "K", "L", "M", "AH"))
    paren = re.compile(r"\([^)]{5,}\)\s*$", re.S)
    fin_bad = []
    for r in data:
        last = re.sub(r"^\s*\d+\.\s*", "", lines(ws.cell(r, 12).value)[-1])
        if "check that" not in last or len(last.split()) > 18:
            fin_bad.append(r)
    checks = {
        "資料列": len(data),
        "Requirement 覆蓋": len({str(ws.cell(r, 4).value).strip() for r in data}),
        "PENDING": sum(1 for r in data if "PENDING" in blob(r)),
        "DR 引用": sum(1 for r in data if "DR-" in blob(r)),
        "BLOCKED": sum(1 for r in data if "BLOCKED" in blob(r)),
        "IMPL_GAP": sum(1 for r in data if "IMPL_GAP" in blob(r)),
        "NOT GENERATED": sum(1 for r in data if "NOT GENERATED" in blob(r)),
        "Proc↔ER 不等": [r for r in data if len(lines(ws.cell(r, 12).value))
                        != len(lines(ws.cell(r, 13).value))],
        "步數<2": [r for r in data if len(lines(ws.cell(r, 12).value)) < 2],
        "尾句號": [r for r in data if any(x.strip().endswith((".", "。"))
                                       for c in A for x in lines(ws.cell(r, C[c]).value))],
        "行首尾空白": [r for r in data if any(
            x != x.strip() for c in A + ("I", "N")
            for x in str(ws.cell(r, C[c]).value or "").split("\n"))],
        "四欄全形標點": [r for r in data if any(
            ch in str(ws.cell(r, C[c]).value or "") for c in A for ch in "；，。、（）")],
        "括號下半缺": [r for r in data
                   if not paren.search(str(ws.cell(r, 9).value or "").strip())],
        "Priority 越域": [r for r in data if str(ws.cell(r, 16).value)
                       not in ("P0", "P1", "P2", "P3")],
        "R-POP20 Final Step": fin_bad,
    }
    print("  3. 全簿現算：")
    for k, v in checks.items():
        print(f"       {k:20} {v}")
    assert checks["資料列"] == 5 and checks["Requirement 覆蓋"] == 5, "§四.3 列數不符"
    assert all(not v for k, v in checks.items()
               if k not in ("資料列", "Requirement 覆蓋")), "§四.3 有違規，停手"

    DELIVERED.mkdir(parents=True, exist_ok=True)
    del_path = DELIVERED / DELIVER_NAME
    shutil.copy2(out_path, del_path)
    s_out, s_del = sha256(out_path), sha256(del_path)
    assert s_out == s_del, "output/ 與 delivered/ 之 sha 不同，停手"
    print(f"  4. output/ 與 delivered/ sha256 相同：{s_out}")

    # ------------------------------------------------ §3.3 MANIFEST
    rel_src = str(SRC.relative_to(ROOT))
    specs = [
        (OUT_DIR / "MANIFEST.tsv",
         ["filename", "sha256", "source_path", "round", "status", "note"],
         [DELIVER_NAME, s_out, rel_src, "PU-02", "delivered", NOTE]),
        (DELIVERED / "MANIFEST.tsv",
         ["filename", "sha256", "source_path", "delivered_round", "note"],
         [DELIVER_NAME, s_del, rel_src, "PU-02", NOTE]),
    ]
    lineno = {}
    for path, header, row in specs:
        fresh = not path.exists()
        if fresh:
            with open(path, "w", encoding="utf-8", newline="") as f:
                csv.writer(f, delimiter="\t", lineterminator="\n").writerow(header)
        n0 = sum(1 for _ in open(path, encoding="utf-8"))
        with open(path, "a", encoding="utf-8", newline="") as f:
            csv.writer(f, delimiter="\t", lineterminator="\n").writerow(row)
        lineno[str(path.relative_to(ROOT))] = (n0 + 1, fresh)

    print("\n== 產出 ==")
    print(f"  surgical_save: {report}")
    print(f"  {out_path.relative_to(ROOT)}")
    print(f"  {del_path.relative_to(ROOT)}")
    for p, (n, fresh) in lineno.items():
        print(f"  MANIFEST {p} 追加第 {n} 行" + ("（本包新建，含表頭）" if fresh else ""))
    print(f"\n  來源 sha256（未動） {sha256(SRC)}")
    print(f"  交付本 sha256       {s_out}")
    assert sha256(SRC)[:16] == SRC_SHA16, "來源 sha 改變，停手"
    return 0


if __name__ == "__main__":
    sys.exit(main())
